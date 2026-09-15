"""Replace only slide 12 with editable dimensional terminals in a triangle."""
from pathlib import Path
from copy import deepcopy
import argparse, zipfile, hashlib, json
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
ROOT=Path(__file__).resolve().parents[3]
p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--output',type=Path,required=True);args=p.parse_args()
assert not args.output.exists(), 'Existing files are protected'
r=Presentation(args.source);s=r.slides[11];original=list(s.shapes)
assert len(r.slides)==36 and original[2].text=='记忆共享，分布互连'
# Retain title, navigation, background, introduction and existing physical page number.
shots=[(deepcopy(original[i]._element),deepcopy(original[i+1]._element)) for i in [24,31,38]]
for sh in original[20:]:s.shapes._spTree.remove(sh._element)
CYAN='4BE5F4';LIGHT='B9F8FF';DARK='05243F'
def shape(kind,x,y,w,h,fill,stroke=None,width=1):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:q.line.color.rgb=RGBColor.from_string(stroke);q.line.width=Pt(width)
 else:q.line.fill.background()
 if kind==MSO_SHAPE.ROUNDED_RECTANGLE:q.adjustments[0]=.08
 return q

def gradient(q,colors):
 sp=q._element.find('p:spPr',NS)
 for el in list(sp):
  if E.QName(el).localname in ['solidFill','gradFill','noFill']:sp.remove(el)
 fill=E.SubElement(sp,'{'+NS['a']+'}gradFill',rotWithShape='1');st=E.SubElement(fill,'{'+NS['a']+'}gsLst')
 for j,c in enumerate(colors):E.SubElement(E.SubElement(st,'{'+NS['a']+'}gs',pos=str(round(j*100000/(len(colors)-1)))),'{'+NS['a']+'}srgbClr',val=c)
 E.SubElement(fill,'{'+NS['a']+'}lin',ang='5400000',scaled='1')
 sp.remove(fill);ln=sp.find('a:ln',NS);sp.insert(list(sp).index(ln) if ln is not None else len(sp),fill)

def poly(points,fill,stroke=None):
 builder=s.shapes.build_freeform(Inches(points[0][0]),Inches(points[0][1]));builder.add_line_segments([(Inches(x),Inches(y)) for x,y in points[1:]],close=True);q=builder.convert_to_shape();q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if stroke:q.line.color.rgb=RGBColor.from_string(stroke);q.line.width=Pt(.8)
 else:q.line.fill.background()
 return q

def label(x,y,w,h,value,size=16,color=LIGHT,bold=False):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 for j,t in enumerate(value.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.alignment=PP_ALIGN.CENTER;qrun=p.add_run();qrun.text=t;qrun.font.name='Microsoft YaHei';qrun.font.size=Pt(size);qrun.font.bold=bold;qrun.font.color.rgb=RGBColor.from_string(color)
 return q

def line(x,y,u,v,width=2,color=CYAN,arrows=False,glow=False):
 q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(u),Inches(v));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(width)
 ln=q._element.find('.//a:ln',NS)
 if arrows:
  for tag in ['headEnd','tailEnd']:E.SubElement(ln,'{'+NS['a']+'}'+tag,type='triangle',w='med',len='med')
 if glow:
  ef=E.SubElement(q._element.find('p:spPr',NS),'{'+NS['a']+'}effectLst');g=E.SubElement(ef,'{'+NS['a']+'}glow',rad='75000');c=E.SubElement(g,'{'+NS['a']+'}srgbClr',val=CYAN);E.SubElement(c,'{'+NS['a']+'}alpha',val='28000')
 return q
# Lower the reading density so the triangle has a clear three-dimensional stage.
intro=original[19];intro.text='可信设备先配对、再授权共享；各端保留本地记忆，离线可用，重连后补齐。'
intro.top=Inches(1.96);intro.height=Inches(.35)
for pa in intro.text_frame.paragraphs:
 pa.alignment=PP_ALIGN.CENTER
 for rr in pa.runs:rr.font.name='Microsoft YaHei';rr.font.size=Pt(16);rr.font.bold=True;rr.font.color.rgb=RGBColor.from_string('FFFFFF')
# Triangular links sit behind the three endpoints. There is no central server.
line(5.06,3.78,3.89,4.94,2.5,arrows=True,glow=True)
line(8.29,3.78,9.47,4.94,2.5,arrows=True,glow=True)
line(4.54,5.86,8.80,5.86,2.5,arrows=True,glow=True)
label(3.30,4.06,1.40,.3,'对等同步',13,CYAN)
label(8.65,4.06,1.40,.3,'对等同步',13,CYAN)
label(4.80,5.42,3.75,.29,'授权共享 · 本地副本',17,CYAN,True)
label(4.80,6.02,3.75,.28,'离线重连后自动补齐',13,LIGHT)
label(5.21,4.91,2.90,.34,'去中心化记忆网络',19,LIGHT,True)

def terminal(x,y,w,shot,kind):
 h=w*9/16; depth=.12
 # Ground shadow and concentric luminous plinth.
 basey=y+h+(.54 if kind!='laptop' else .38)
 sh=shape(MSO_SHAPE.OVAL,x-.30,basey-.12,w+.6,.34,'031B35')
 sh=shape(MSO_SHAPE.OVAL,x-.23,basey-.10,w+.46,.24,'074563',CYAN,.6)
 sh=shape(MSO_SHAPE.OVAL,x-.12,basey-.08,w+.24,.15,'063550')
 # Back and side surfaces give the chassis explicit thickness.
 shape(MSO_SHAPE.ROUNDED_RECTANGLE,x+.12,y-.10,w+.20,h+.25,'0B304C','3F8098',.7)
 poly([(x+w+.10,y),(x+w+.23,y-.08),(x+w+.23,y+h+.12),(x+w+.10,y+h+.23)],'0C2942','478DA5')
 if kind!='laptop':
  poly([(x+w*.44,y+h+.15),(x+w*.59,y+h+.15),(x+w*.64,basey-.07),(x+w*.40,basey-.07)],'5295A8','95DBE8')
  poly([(x+w*.25,basey-.07),(x+w*.66,basey-.13),(x+w*.79,basey+.02),(x+w*.30,basey+.07)],'458299','9ADAE9')
  line(x+w*.32,basey+.035,x+w*.72,basey+.005,.8,LIGHT)
 bezel=shape(MSO_SHAPE.ROUNDED_RECTANGLE,x-.06,y-.04,w+.18,h+.25,'13384F',LIGHT,.8);gradient(bezel,['7CB9C8','21455B','061B2C'])
 shape(MSO_SHAPE.RECTANGLE,x+.01,y+.03,w+.04,h+.04,'030D1A')
 # Original, full desktop capture + proportionally positioned editable red rectangle.
 pic,mark=deepcopy(shot[0]),deepcopy(shot[1]);tree=s.shapes._spTree
 # Unique IDs for the reused native screenshot shapes.
 nextid=max(int(z.get('id')) for z in tree.xpath('.//p:cNvPr'))+1
 for j,el in enumerate([pic,mark]):el.find('.//p:cNvPr',NS).set('id',str(nextid+j));tree.append(el)
 im=s.shapes[-2];ann=s.shapes[-1];ox,oy=im.left,im.top;factor=Inches(w)/im.width
 for q in [im,ann]:
  q.left=Inches(x+.03)+round((q.left-ox)*factor);q.top=Inches(y+.05)+round((q.top-oy)*factor);q.width=round(q.width*factor);q.height=round(q.height*factor)
 label(x+.12,y+h+.07,w-.24,.11,'P I X I U',6,LIGHT,True)
 shape(MSO_SHAPE.OVAL,x+w-.10,y+h+.105,.035,.035,CYAN)
 if kind=='desktop':
  # Separate workstation tower.
  tx=x+w+.34;ty=y+.39
  poly([(tx,ty),(tx+.17,ty-.13),(tx+.55,ty-.13),(tx+.38,ty)],'7296A8','B4DCE7')
  poly([(tx+.38,ty),(tx+.55,ty-.13),(tx+.55,ty+1.28),(tx+.38,ty+1.41)],'0B2C44','4E91AD')
  q=shape(MSO_SHAPE.ROUNDED_RECTANGLE,tx,ty,.38,1.41,'244C63','86C6DB');gradient(q,['4D829A','0A253D'])
  shape(MSO_SHAPE.OVAL,tx+.145,ty+.14,.085,.085,CYAN)
  for j in range(5):line(tx+.09,ty+.60+j*.065,tx+.29,ty+.60+j*.065,.7,'77B3C6')
 elif kind=='laptop':
  # Perspective keyboard deck, front lip and trackpad.
  top=y+h+.19
  poly([(x-.02,top),(x+w+.12,top),(x+w+.40,top+.41),(x-.31,top+.41)],'53859D',LIGHT)
  poly([(x-.31,top+.41),(x+w+.40,top+.41),(x+w+.29,top+.49),(x-.22,top+.49)],'183F58','66B7CE')
  for row in range(3):
   for col in range(11):
    xx=x+.19+col*(w-.28)/11-row*.02;yy=top+.065+row*.065
    poly([(xx,yy),(xx+.18,yy),(xx+.19,yy+.037),(xx-.01,yy+.037)],'142D44')
  poly([(x+w*.39,top+.275),(x+w*.62,top+.275),(x+w*.65,top+.36),(x+w*.37,top+.36)],'88B3C4','B2DBE7')
 return basey
# Top vertex: workstation. Lower vertices: all-in-one and travel laptop.
terminal(5.15,2.72,2.60,shots[0],'desktop')
label(4.99,2.34,3.35,.30,'设备 A · 书房台式机',17,CYAN,True)
terminal(1.18,4.63,2.75,shots[1],'allinone')
label(.68,4.23,3.80,.31,'设备 B · 客厅一体机',18,CYAN,True)
terminal(9.36,4.63,2.75,shots[2],'laptop')
label(8.87,4.23,3.87,.31,'设备 C · 随身笔记本',18,CYAN,True)
label(4.62,6.59,4.14,.25,'同一条记忆，三端持续接续',16,CYAN,True)
label(2.90,6.95,7.55,.30,'11 月 10 日下午三点 · 三楼档案室',20,'FFD29C',True)
with zipfile.ZipFile(args.source) as z:blobs={n:z.read(n) for n in z.namelist()}
part=str(s.part.partname).lstrip('/');blobs[part]=E.tostring(s._element,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(args.output,'w',zipfile.ZIP_DEFLATED) as z:
 for n,data in blobs.items():z.writestr(n,data)
print(args.output)

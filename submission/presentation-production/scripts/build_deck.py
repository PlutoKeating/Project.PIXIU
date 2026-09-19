"""Editable video-aligned presentation: every claim maps to the delivered film."""
from pathlib import Path
import json, hashlib, math
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE as S, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
R=Path(__file__).resolve().parents[3];W=R/'submission/presentation-production';V=R/'submission/video-production';A=R/'docs/delivery/assets/operations/03-current-workflows'
B='1456B8';D='102A50';I='182E4D';M='58708C';P='EAF2FF';C='37B6D3';L='CEDFF2';F='FFFFFF';FONT='Noto Sans CJK SC'
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
prs.core_properties.title='PIXIU·貔貅｜个人记忆助手';prs.core_properties.author='PIXIU';prs.core_properties.last_modified_by='PIXIU'
timeline=json.loads((V/'src/timeline.json').read_text());shots={s['id']:s for s in timeline['shots']};records=[];toc_links=[]
inputs={Path(__file__).resolve(),W/'source/video-aligned-plan.md',V/'src/timeline.json',V/'src/directed-scenes.json',V/'review/directed-final-package.json',V/'src/sync-trace.json',R/'build/release/scripts/build-presentation.py',R/'build/release/requirements-docs.txt',W/'requirements.txt'}
def el(tag,**attrs):
 e=OxmlElement(tag)
 for k,v in attrs.items():e.set(k,str(v))
 return e
def shape(s,k,x,y,w,h,fill=F,stroke=None,sw=1):
 a=s.shapes.add_shape(k,Inches(x),Inches(y),Inches(w),Inches(h))
 if k==S.ROUNDED_RECTANGLE:a.adjustments[0]=.08
 if fill:a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(fill)
 else:a.fill.background()
 if stroke:a.line.color.rgb=RGBColor.from_string(stroke);a.line.width=Pt(sw)
 else:a.line.fill.background()
 a._element.spPr.append(el('a:effectLst'));return a
def rect(s,x,y,w,h,fill=F,stroke=None):return shape(s,S.ROUNDED_RECTANGLE,x,y,w,h,fill,stroke)
def circle(s,x,y,d,fill=P,stroke=None,sw=1):return shape(s,S.OVAL,x,y,d,d,fill,stroke,sw)
def txt(s,t,x,y,w,h,size=20,col=I,bold=False,align=None):
 a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=a.text_frame;tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
 for j,row in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.text=row;p.font.name=FONT;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(col);p.space_after=Pt(6);p.line_spacing=1.08
  if align is not None:p.alignment=align
  for run in p.runs:run._r.get_or_add_rPr().append(el('a:ea',typeface=FONT))
 return a
def line(s,x,y,x2,y2,col=B,width=2,arrow=False,dash=False):
 a=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(x2),Inches(y2));a.line.color.rgb=RGBColor.from_string(col);a.line.width=Pt(width);a._element.spPr.append(el('a:effectLst'));ln=a.line._get_or_add_ln()
 if dash:ln.append(el('a:prstDash',val='dash'))
 if arrow:ln.append(el('a:tailEnd',type='triangle'))
 return a
def grad(a,c1,c2):
 sp=a._element.spPr
 for e in list(sp):
  if e.tag.split('}')[-1] in ('solidFill','noFill','gradFill'):sp.remove(e)
 g=el('a:gradFill',rotWithShape='1');ls=el('a:gsLst')
 for pos,col in [(0,c1),(100000,c2)]:
  st=el('a:gs',pos=pos);st.append(el('a:srgbClr',val=col));ls.append(st)
 g.append(ls);g.append(el('a:lin',ang=2700000,scaled='1'))
 # Fill precedes line/effects per DrawingML schema.
 sp.insert(next((i for i,e in enumerate(sp) if e.tag.split('}')[-1] in ('ln','effectLst','effectDag')),len(sp)),g)
def badge(s,t,x,y,w=1.55,dark=False):
 rect(s,x,y,w,.38,'244E80' if dark else P);txt(s,t,x+.1,y+.065,w-.2,.23,12,'BDDFF8' if dark else B,True)
def num(s,n,x,y,d=.5,dark=False):
 circle(s,x,y,d,C if dark else B);txt(s,str(n),x,y+.1,d,.29,16,D if dark else F,True,PP_ALIGN.CENTER)
def timecode(f):
 t=round(f/30);return f'{t//60:02}:{t%60:02}'
def page(title,chapter,ids,layout,dark=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]);bg=shape(s,S.RECTANGLE,0,0,13.333333,7.5,F)
 if dark:grad(bg,D,'1C518B')
 rec={'page':len(prs.slides),'title':title,'chapter':chapter,'layout':layout,'shots':ids,'screenshots':[]};records.append(rec)
 txt(s,'PIXIU / 貔貅',.55,.30,2.4,.32,14,F if dark else B,True)
 txt(s,chapter,8.0,.32,4.72,.3,12,'B6CEE8' if dark else M,align=PP_ALIGN.RIGHT)
 if ids:
  start=min(shots[k]['from'] for k in ids);end=max(shots[k]['from']+shots[k]['duration'] for k in ids);rec['video_range']=[timecode(start),timecode(end)]
  txt(s,'宣传片 '+timecode(start)+'–'+timecode(end),.55,7.12,4,.22,10,'B6CEE8' if dark else M)
 txt(s,f'{len(prs.slides):02}',12.12,7.04,.6,.37,14,'B6CEE8' if dark else B,True,PP_ALIGN.RIGHT)
 s.notes_slide.notes_text_frame.text=title+'\n内容依据：正式演示增强版499.333秒；以下为对应镜头原口白。\n'+'\n\n'.join(k+' '+shots[k]['title']+'\n'+shots[k]['narration'] for k in ids)+'\n实拍为0.1.12银河麒麟V11公开合成演示资料。'
 return s
def heading(s,title,sub=None,dark=False):
 txt(s,title,.66,1.02,12,.76,34,F if dark else I,True)
 if sub:txt(s,sub,.69,1.99,11.9,.54,18,'C7DDF2' if dark else M)
def pic(s,name,x,y,w,h,caption=None):
 p=Path(name) if isinstance(name,Path) else A/(name+'.png');inputs.add(p)
 im=Image.open(p);iw,ih=im.size;scale=min(w/iw,h/ih);dw,dh=iw*scale,ih*scale;px=x+(w-dw)/2;py=y+(h-dh)/2
 rect(s,px-.035,py-.035,dw+.07,dh+.07,F,L)
 s.shapes.add_picture(str(p),Inches(px),Inches(py),Inches(dw),Inches(dh))
 records[-1]['screenshots'].append({'path':p.relative_to(R).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()})
 if caption:txt(s,caption,x,y+h+.1,w,.34,12,M)
def film(s,sid,x,y,w,h,phase=72):pic(s,W/f'assets/video-frames/{sid}-{phase}.png',x,y,w,h)
def label(s,title,body,x,y,w=3.4,col=B):
 txt(s,title,x,y,w,.42,22,col,True);txt(s,body,x,y+.63,w,1.15,18,I)
def monitor(s,x,y,w=2.8,title='本地记忆',dark=False):
 col='7EC9EB' if dark else B;rect(s,x,y,w,w*.62,'214B7B' if dark else P,col)
 rect(s,x+.1,y+.1,w-.2,w*.43,'183A65' if dark else F)
 txt(s,title,x+.18,y+w*.18,w-.36,.4,18,F if dark else B,True,PP_ALIGN.CENTER)
 line(s,x+w*.5,y+w*.62,x+w*.5,y+w*.76,col,3);line(s,x+w*.25,y+w*.76,x+w*.75,y+w*.76,col,3)
def chapter(n,title,sub,ids,icon):
 s=page(title,f'第{n}章',ids,'chapter-'+icon,True)
 txt(s,f'{n:02}',.65,1.23,3.3,2.1,106,'4376AA',True)
 txt(s,title,.73,3.56,7.7,1.45,39,F,True);txt(s,sub,.77,5.56,7.5,.9,20,'C5DFF6')
 if icon=='files':
  for j,(t,xx,yy) in enumerate([('任务',8.45,2.0),('资料',9.5,2.75),('来源',10.53,3.5)]):
   shape(s,S.FOLDED_CORNER,xx,yy,1.76,2.2,'234E80','73B8DF',1.4);txt(s,t,xx+.3,yy+1.1,1.2,.4,23,F,True)
 elif icon=='orbit':
  for d in [3.0,4.0]:circle(s,10-d/2,3.65-d/2,d,None,'5185B7')
  for k,t in enumerate(['对话','工具','记录','行为']):
   x=9.55+1.6*math.cos(k*math.pi/2);y=3.25+1.6*math.sin(k*math.pi/2);circle(s,x,y,.95,'245887','6FB8DE');txt(s,t,x,y+.3,.95,.35,17,F,True,PP_ALIGN.CENTER)
 elif icon=='devices':
  for x,y in [(8.1,1.8),(10.4,3.6),(7.9,5.0)]:monitor(s,x,y,1.8,'记忆',True)
  line(s,9.9,2.8,10.7,3.7,C,2,True);line(s,10.45,4.7,9.7,5.4,C,2,True);line(s,8.8,4.9,8.8,3.35,C,2,True)
 elif icon=='scope':
  for x,y,d in [(8,1.9,4.3),(8.55,2.45,3.2),(9.13,3.03,2.04)]:circle(s,x,y,d,None,'6AB4DE',2)
  txt(s,'用户\n掌握边界',9.21,3.59,1.86,1.03,22,F,True,PP_ALIGN.CENTER)
 elif icon=='layers':
  for j,t in enumerate(['智能体宿主','记忆适配','知识与同步']):
   rect(s,8.3+j*.2,2.0+j*1.33,3.58,1.04,'255588','79BCE2');txt(s,t,8.55+j*.2,2.29+j*1.33,3.08,.4,22,F,True)
 else:
  for j,(t,v) in enumerate([('偏好','准确'),('检索','召回'),('响应','时延'),('更正','正确')]):
   x=8.3+(j%2)*2.05;y=2+(j//2)*2.1;circle(s,x,y,1.7,None,'66B5DD',2);txt(s,t,x,y+.32,1.7,.35,18,'A7DDF2',align=PP_ALIGN.CENTER);txt(s,v,x,y+.82,1.7,.45,24,F,True,PP_ALIGN.CENTER)
 return s
# 01 — film identity, not an invented proposition.
s=page('PIXIU·貔貅','项目报告',['s01'],'brand-stage',True)
txt(s,'PIXIU',.7,1.25,6.1,1.2,72,F,True);txt(s,'貔貅',.79,2.65,4,.9,43,F,True)
txt(s,'面向麒麟操作系统智能体的\n去中心化记忆系统',.79,4.02,6.4,1.14,27,'CAE4F8')
txt(s,'让每一台设备的记忆，彼此相通。',.81,6.18,10.8,.6,26,F,True)
for n,(a,x,y,w,h) in enumerate([('shared-workspace',7.55,1.2,4.9,3.1),('directory-recall',8.45,3.8,3.95,2.0),('preference-history',6.72,4.16,2.32,1.4)]):pic(s,a,x,y,w,h)
# 02 — explicit TOC, actual pages and internal links.
s=page('目录','内容导航',[],'editorial-index')
txt(s,'目\n录',.7,1.38,2.2,3.7,70,B,True);txt(s,'沿着一次使用，\n理解记忆如何积累、\n复用和共享。',.78,5.02,3.8,1.3,21,M)
chapters=[('01','日常资料，持续积累','任务执行 / 目录整理 / 账单追溯',3),('02','让经验留下来','多源接入 / 偏好 / 版本 / 分层记忆',10),('03','可信设备，共同记忆','配对 / 跨端使用 / 离线与并发',18),('04','每份记忆，都有边界','读取与保存范围 / 信任 / 精准遗忘',23),('05','体验背后，如何协同','系统分工 / 联合检索 / 安装更新',26),('06','用结果，检验价值','量化测试 / 洞察简报 / 应用价值',30)]
for j,(n,t,b,p) in enumerate(chapters):
 y=1.29+j*.91;num(s,n,4.68,y,.5);link=txt(s,t,5.47,y-.01,5.9,.38,23,I,True);txt(s,b,5.48,y+.45,6.15,.29,13,M);txt(s,f'{p:02}',12.05,y,.56,.43,24,B,True,PP_ALIGN.RIGHT);toc_links.append((link,p));line(s,5.47,y+.78,12.63,y+.78,L,.7)
# 03
chapter(1,'日常资料，\n持续积累','从一次任务开始，让过去的安排在下一次使用中继续发挥作用。',['s02'],'files')
# 04
s=page('两件事，让日常经验继续发挥作用','01 / 使用价值',['s02','s03'],'two-worlds')
rect(s,0,3.0,6.55,3.55,P);rect(s,6.77,3.0,6.57,3.55,'F2F7FC');heading(s,'把资料留下，把经验接续')
txt(s,'01',.78,2.12,1.0,.7,40,B,True);txt(s,'持续积累成知识',1.85,2.25,4.5,.45,26,I,True)
for j,t in enumerate(['保存资料','查找安排','带着来源']):
 shape(s,S.FOLDED_CORNER,1.0+j*1.73,3.45,1.22,1.5,F,L);txt(s,t,1.0+j*1.73,5.2,1.6,.39,18,B,True)
txt(s,'从日常内容，整理出可查询的记忆。',.94,6.0,5.3,.4,20,I)
txt(s,'02',7.08,2.12,1,.7,40,B,True);txt(s,'可信设备共同使用',8.15,2.25,4.6,.45,26,I,True)
for x,y in [(7.25,3.5),(10.15,3.5),(8.7,4.9)]:monitor(s,x,y,1.75,'本地副本')
line(s,9.04,4,10.07,4,C,2,True);line(s,10.5,4.91,10.2,5.3,C,2,True);line(s,8.75,5.3,8.2,4.9,C,2,True)
txt(s,'经过授权，彼此交换更新。',7.13,6.39,5.4,.4,20,I)
# 05
s=page('一个应用，四个清晰入口','01 / 原生桌面',['s04'],'product-spotlight');heading(s,'一个应用，四个清晰入口')
pic(s,'shared-workspace',2.82,2.25,7.73,4.4)
for t,b,x,y in [('会话','交代任务',.77,2.65),('记忆','管理已保存内容',.77,4.6),('设备','连接协作电脑',10.78,2.65),('设置','管理授权',10.78,4.6)]:label(s,t,b,x,y,1.95)
badge(s,'银河麒麟 V11',4.02,1.84,2.13);badge(s,'系统文本向量化 / 向量数据库',6.33,1.84,4.33)
# 06
s=page('从采购任务，到可复用的结果','01 / 任务执行',['s06'],'giant-number-ledger')
txt(s,'从采购任务，\n到可复用的结果',.7,1.08,5.4,1.6,34,I,True)
txt(s,'496',.76,2.96,4.4,1.23,74,B,True);txt(s,'元',4.87,3.6,.8,.5,26,M)
for j,(a,b) in enumerate([('12套书籍 × 38.50','462.00'),('5包标签纸 × 6.80','34.00')]):
 y=4.5+j*.62;txt(s,a,.82,y,3.8,.36,20,I);txt(s,b,4.4,y,1.37,.36,20,B,True,PP_ALIGN.RIGHT)
line(s,.8,5.8,5.77,5.8,L,1);txt(s,'核算明细 → 写入文件 → 保存记忆',.84,6.1,5.25,.5,20,B,True)
pic(s,'agent-tools-completed',6.36,1.62,6.12,2.29,'真实任务结果：计算、文件与记忆');pic(s,'agent-task-recall',6.36,4.47,6.12,1.86,'打开新会话，继续查询数量和金额')
# 07
s=page('保存一份资料，后台自动整理','01 / 目录整理',['s11'],'vertical-journey');heading(s,'保存一份资料，后台自动整理','社区筹备星河观测活动：资料进入授权目录，下一次直接询问安排。')
for j,(t,b) in enumerate([('授权目录','开启目录采集并保存'),('放入资料','后台整理，显示处理进度'),('新会话查询','找回集合时间、地点与预算'),('点击来源','阅读活动原文')]):
 y=2.9+j*.84;num(s,j+1,.86,y);txt(s,t,1.63,y,3.42,.39,23,B,True);txt(s,b,1.65,y+.44,3.6,.33,16,M)
 if j<3:line(s,1.11,y+.52,1.11,y+.79,L,1.5)
pic(s,'directory-recall',5.66,2.89,3.52,3.35);pic(s,'directory-source',9.56,2.89,2.88,3.35)
txt(s,'找到安排',6.11,6.38,2.6,.4,19,B,True);txt(s,'核对原文',9.86,6.38,2.4,.4,19,B,True)
# 08
s=page('资料更新，先对照，再保存','01 / Dreaming 更正',['s11b'],'dark-before-after',True);heading(s,'资料更新，先对照，再保存',dark=True)
pic(s,'dreaming-review',.75,2.3,6.42,4.26)
txt(s,'后台阅读通知，提出更正建议',7.66,2.31,4.91,.57,23,'C9E4F8',True)
txt(s,'原安排',7.68,3.26,4.5,.33,16,'B9D5EF');txt(s,'11月8日 19:00 · 三层',7.68,3.75,4.74,.46,25,F)
line(s,9.93,4.41,9.93,4.8,C,2.4,True)
txt(s,'核对并批准后',7.68,4.99,4.5,.33,16,'B9D5EF');txt(s,'11月15日 19:30\n社区天文台二层',7.68,5.47,4.98,1.03,29,F,True)
badge(s,'预算保持 2680元',7.67,6.56,3.15,True)
# 09
s=page('新会话查账，沿来源复核','01 / 跨会话与来源',['s07','s08'],'bill-evidence-spread')
txt(s,'新会话查账，\n沿来源复核',.71,1.08,5.15,1.55,35,I,True)
txt(s,'434.50',.77,3.05,5.4,1.1,60,B,True);txt(s,'元 · 九月水电燃气合计',.84,4.29,4.9,.45,22,M)
for j,(a,b) in enumerate([('电费','210.00'),('水费','68.50'),('燃气费','156.00')]):
 y=5.1+j*.48;txt(s,a,.84,y,2.1,.35,19,I);txt(s,b,3.45,y,1.85,.35,19,B,True,PP_ALIGN.RIGHT)
pic(s,'bill-recall',6.3,1.32,3.83,4.23);pic(s,'bill-source',9.44,3.66,3.05,2.65)
txt(s,'提出问题 → 找到账单 → 点击来源',6.45,6.59,6.0,.39,20,B,True)
# 10
chapter(2,'让经验\n留下来','对话、工具、记录与习惯，逐步成为可以重复使用的知识。',['s09'],'orbit')
# 11
s=page('多种来源，进入统一记忆','02 / 多源接入',['s09'],'radial-ingest');heading(s,'多种来源，进入统一记忆')
for j,(t,x,y) in enumerate([('日常对话',1,2.53),('工具结果',1,4.93),('应用使用',9.65,2.53),('手动录入',9.65,4.93)]):
 circle(s,x,y,2.13,P,B,1.2);txt(s,t,x+.15,y+.79,1.83,.45,23,B,True,PP_ALIGN.CENTER)
 line(s,x+(2.18 if x<5 else -.08),y+1.06,5.09 if x<5 else 8.23,4.22,C,2,True)
circle(s,4.99,2.56,3.35,None,L,1.6);circle(s,5.21,2.78,2.91,B);txt(s,'统一整理\n保留来源',5.58,3.7,2.2,1.02,29,F,True,PP_ALIGN.CENTER)
for j,t in enumerate(['格式检查','质量检查','敏感检查']):badge(s,t,4.54+j*1.48,6.35,1.36)
# 12
s=page('明确的约定，手动记下来','02 / 手动记录',['s10'],'form-with-annotations');heading(s,'明确的约定，手动记下来','填写标题和正文，决定保留范围；保存后立即查询核对。')
pic(s,'manual-entry',.78,2.74,7.48,3.87)
for j,(t,b) in enumerate([('写清内容','标题 + 正文'),('选择范围','个人或共享'),('保存后核对','查询刚录入的事项')]):
 y=2.9+j*1.17;num(s,j+1,8.77,y);txt(s,t,9.52,y-.02,2.87,.42,23,B,True);txt(s,b,9.54,y+.53,2.88,.44,19,M)
# 13
s=page('使用习惯，按授权记录','02 / 行为采集',['s12'],'behavior-collage')
txt(s,'使用习惯，\n按授权记录',.74,1.1,5.2,1.6,35,I,True)
pic(s,'behavior-demo-window',.8,3.04,5.08,2.65);pic(s,'behavior-source-body',6.31,2.26,6.07,3.23)
badge(s,'应用窗口标题',6.36,1.51,2.34);badge(s,'使用时长',8.94,1.51,1.85)
txt(s,'28',6.4,5.75,1.8,.85,52,B,True);txt(s,'秒 · 片中这次使用记录',8.05,6.12,4.4,.4,20,M)
txt(s,'先开启授权，再查看对应来源。\n采集权限由设置页统一管理。',.88,6.0,5.15,.79,18,B,True)
# 14
s=page('习惯变化，记忆跟着更新','02 / 偏好记忆',['s15'],'preference-timeline');heading(s,'习惯变化，记忆跟着更新','从“回答简洁”到“解释充分”，保存当前偏好，也保留变化历史。')
pic(s,'preference-history',.79,2.72,6.98,3.84)
line(s,8.77,3.12,8.77,6.29,L,3)
for j,(v,t,b) in enumerate([('v1','简洁回答','日常问答'),('v2','仍然简洁','历史更新保留'),('v3','详细说明','后续会话使用新偏好')]):
 y=2.9+j*1.21;num(s,v,8.5,y,.55);txt(s,t,9.37,y,2.99,.42,24,B,True);txt(s,b,9.4,y+.55,3,.42,17,M)
# 15
s=page('把工作经验，整理成四类知识','02 / 知识复用',['s14'],'knowledge-books');heading(s,'把工作经验，整理成四类知识','社区图书角：从书架位置到归还流程，为下一次相似任务留下经验。')
for j,(t,b) in enumerate([('事实','书架位置'),('流程','图书归还'),('案例','分类调整'),('模板','阅读计划')]):
 x=.81+j*3.14;rect(s,x,2.8,2.81,1.8,[P,'DDEBF9','D3E4F6','C7DEF4'][j],L);shape(s,S.FOLDED_CORNER,x+2.18,3.02,.34,.4,F,None);txt(s,t,x+.22,3.07,2.2,.48,27,B,True);txt(s,b,x+.23,3.8,2.2,.42,20,I)
pic(s,'knowledge-workflow-body',.91,5.07,6.39,1.51);txt(s,'① 核对书名\n② 按主题分类、上架\n③ 更新借阅登记',7.7,5.0,4.74,1.63,22,B,True)
# 16
s=page('更新有版本，冲突有记录','02 / 维护知识',['s16'],'version-comparison');heading(s,'更新有版本，冲突有记录')
txt(s,'周五 16:00',1.06,2.19,4.57,.73,40,M,True);line(s,5.74,2.63,7.25,2.63,C,2.5,True);txt(s,'周五 17:00',7.7,2.19,4.66,.73,40,B,True)
pic(s,'edit-version-one',.91,3.47,5.45,2.53,'原记录');pic(s,'edit-version-two',7.02,3.47,5.45,2.53,'保存后重新读取：版本二')
txt(s,'需要人工选择时，到冲突页对照候选内容与来源。',1.0,6.57,11.6,.39,22,B,True)
# 17
s=page('短期关注，中期整理，长期复用','02 / 记忆流转',['s17'],'tier-staircase');heading(s,'短期关注，中期整理，长期复用')
for j,(t,b) in enumerate([('短期','当前对话中的要点'),('中期','跨会话的阶段进展'),('长期','反复使用的知识')]):
 x=.83+j*2.58;y=4.8-j*.8;rect(s,x,y,2.32,1.51,[P,'D7E8F9',B][j]);txt(s,t,x+.22,y+.24,1.88,.48,28,F if j==2 else B,True);txt(s,b,x+.22,y+.94,1.9,.38,15,F if j==2 else I)
 if j<2:line(s,x+2.33,y+.38,x+2.57,y-.38,C,2,True)
film(s,'s17',8.53,2.63,3.96,3.4);txt(s,'打开阶段记录\n选择长期保留\n后续检索与引用',8.65,5.76,3.87,1.06,19,B,True)
# 18
chapter(3,'可信设备，\n共同记忆','换一台电脑，接着使用已经保存的安排与经验。',['s18'],'devices')
# 19
s=page('先建立信任，再共享记忆','03 / 设备配对',['s19'],'pairing-cinema');heading(s,'先建立信任，再共享记忆')
film(s,'s19',.8,2.29,8.22,4.36)
for j,(t,b) in enumerate([('打开配对','确认协作电脑'),('交换信息','建立设备信任'),('查看列表','检查连接和同步状态')]):
 y=2.64+j*1.3;num(s,j+1,9.46,y);txt(s,t,10.14,y,2.33,.44,23,B,True);txt(s,b,9.5,y+.59,3.06,.57,18,M)
# 20
s=page('这里记住，那里接着使用','03 / 跨端使用',['s20'],'two-device-bridge');heading(s,'这里记住，那里接着使用','家庭共享空间中的同一份约定，从书房接续到客厅。')
monitor(s,.88,2.98,3.58,'书房工作站');monitor(s,8.88,2.98,3.58,'客厅一体机');line(s,4.71,3.84,8.58,3.84,C,3,True)
txt(s,'周六 9:00',4.83,2.73,3.56,.66,34,B,True,PP_ALIGN.CENTER);txt(s,'整理书房\n按主题分类书籍\n更新借阅登记',4.95,4.28,3.3,1.27,21,I,align=PP_ALIGN.CENTER)
txt(s,'选择家庭共享空间，保存安排',.89,6.15,4.03,.7,20,B,True);txt(s,'新会话查询，并核对同一来源',8.59,6.15,4.0,.7,20,B,True)
pic(s,'shared-recall',8.97,3.09,3.38,1.41);txt(s,'客厅一体机',9.11,4.71,3.1,.39,18,B,True,PP_ALIGN.CENTER)
# 21
s=page('离线仍可用，重连再对账','03 / 离线协作',['s21'],'offline-swimlane');heading(s,'离线仍可用，重连再对账','资料移交约定：暂停同步期间使用本地记忆；恢复连接后交换更新。')
for j,(t,b) in enumerate([('连接时','三端保存同一份约定'),('暂时离线','本地使用，并保存修改'),('恢复连接','交换更新，补齐差异')]):
 x=2.54+j*3.34;badge(s,t,x,2.77,2.76);txt(s,b,x,3.39,2.9,.53,17,M)
for j,t in enumerate(['设备 A','设备 B','设备 C']):
 y=4.25+j*.8;txt(s,t,.8,y-.14,1.45,.42,20,B,True);line(s,2.4,y+.12,12.43,y+.12,L,2)
 for k in range(3):circle(s,3.6+k*3.34,y-.06,.35,B if k!=1 or j!=2 else F,B,1.5)
line(s,7.12,4.37,10.17,5.97,C,2,True);line(s,7.12,5.17,10.17,5.97,C,2,True)
txt(s,'每台电脑保留本地副本；重连后继续协作。',2.69,6.67,9.43,.35,20,B,True)
# 22
s=page('同时修改，也能收敛','03 / 并发更正',['s22'],'branch-convergence',True);heading(s,'同时修改，也能收敛','记录修改关系 → 按统一规则选择 → 同步到各台设备',True)
for t,b,x in [('设备 A · v2','11/9 10:00\n二楼档案室',.88),('设备 B · v2','11/10 15:00\n三楼档案室',5.08)]:
 rect(s,x,2.91,3.65,1.79,'244F7F','6B9FCC');txt(s,t,x+.25,3.17,3.15,.47,24,F,True);txt(s,b,x+.25,3.85,3.14,.7,21,'D5E8F8')
line(s,2.72,4.81,6.58,5.42,C,2,True);line(s,6.92,4.81,6.58,5.42,C,2,True)
rect(s,3.0,5.56,7.4,1.02,F);txt(s,'三端一致：正文、版本与来源',3.3,5.86,6.8,.49,27,B,True,PP_ALIGN.CENTER)
txt(s,'统一规则',9.3,3.48,3.11,.64,30,F,True);txt(s,'最终：v2\n11/10 15:00\n三楼档案室',9.4,4.29,3.02,1.13,20,'CAE2F6')
# 23
chapter(4,'每份记忆，\n都有边界','决定助手能读什么、记到哪里，也决定何时停止使用。',['s23'],'scope')
# 24
s=page('读取、保存与信任，分别管理','04 / 范围控制',['s23'],'scope-settings');heading(s,'读取、保存与信任，分别管理')
pic(s,'shared-settings',.85,2.49,7.31,3.77)
for j,(t,b) in enumerate([('个人记忆','保存在本机'),('共享内容','用户选择协作设备'),('设置核对','读取范围与新记忆保存位置')]):
 y=2.5+j*1.3;circle(s,8.8,y,.36,B);txt(s,t,9.46,y-.07,2.91,.49,25,B,True);txt(s,b,8.8,y+.57,3.78,.56,18,M)
txt(s,'到设备页管理信任关系，敏感检查保护资料使用范围。',.95,6.6,11.4,.45,21,B,True)
# 25
s=page('先看影响，再确认遗忘','04 / 精准遗忘',['s24'],'forget-sequence');heading(s,'先看影响，再确认遗忘','临时事项结束后，先核对目标和范围，再执行确认。')
for j,(t,b) in enumerate([('提出指令','忘记已完成的演示事项'),('查看预览','核对目标和影响范围'),('点击确认','更新状态，清理检索向量')]):
 x=.87+j*4.16;num(s,j+1,x,2.92);txt(s,t,x+.74,2.91,3.0,.43,25,B,True);txt(s,b,x,3.65,3.8,.54,19,M)
 if j<2:line(s,x+3.64,3.15,x+4.01,3.15,C,2,True)
film(s,'s24',.91,4.52,5.35,2.0);txt(s,'共享记忆的遗忘状态\n也会到达其他设备',7.0,4.77,5.32,1.03,29,B,True);txt(s,'离线设备在重连后补齐。',7.04,6.08,5.1,.43,21,M)
# 26
chapter(5,'体验背后，\n如何协同','宿主、运行时、记忆适配与系统能力，通过明确接口完成协作。',['s05'],'layers')
# 27
s=page('记忆能力，接入完整智能体','05 / 系统分工',['s05'],'layered-architecture');heading(s,'记忆能力，接入完整智能体')
rect(s,.85,2.4,11.64,.91,P);txt(s,'openKylin 宿主与运行时',1.13,2.62,5.75,.43,25,B,True);txt(s,'会话 · 任务规划 · 工具',7.1,2.65,4.97,.41,21,I)
line(s,6.65,3.38,6.65,3.84,C,2.4,True);rect(s,4.76,3.93,3.82,.7,B);txt(s,'PIXIU 记忆适配层',5.03,4.1,3.29,.4,24,F,True,PP_ALIGN.CENTER)
line(s,6.65,4.69,6.65,4.9,C,2);line(s,2.29,4.9,11.71,4.9,C,2)
for xx in [2.29,5.43,8.57,11.71]:line(s,xx,4.9,xx,5.15,C,2,True)
for j,t in enumerate(['资料接入','知识检索','偏好与安全','设备同步']):
 x=.88+j*3.14;rect(s,x,5.22,2.82,1.03,P,L);txt(s,t,x+.12,5.55,2.57,.46,23,B,True,PP_ALIGN.CENTER)
txt(s,'PIXIU 引擎连接记忆业务；系统能力通过接口协作。',1.0,6.64,11.45,.45,22,I,True)
# 28
s=page('关键词、语义、关系，联合找回','05 / 联合检索',['s13'],'query-lens');heading(s,'关键词、语义、关系，联合找回')
txt(s,'“望远镜”\n“器材预算”',.79,2.56,5.33,1.56,39,B,True)
for j,t in enumerate(['关键词','意思相近的表达','关联信息']):badge(s,t,.91,4.58+j*.6,3.96)
pic(s,'keyword-result',6.35,2.52,6.02,1.55);pic(s,'keyword-source-body',6.35,4.62,6.02,1.64)
line(s,5.1,5.51,6.06,5.51,C,2.4,True);txt(s,'找到观测活动 → 打开来源 → 阅读安排与步骤',6.38,6.54,6.03,.47,18,B,True)
# 29
s=page('一个安装包，统一管理','05 / 部署维护',['s27'],'package-and-maintenance');heading(s,'一个安装包，统一管理')
shape(s,S.CUBE,.99,2.52,3.35,2.96,P,B,1.5);txt(s,'PIXIU',1.62,3.63,2.33,.62,35,B,True)
for j,t in enumerate(['桌面应用','助手运行环境','记忆服务']):badge(s,t,4.79,2.92+j*.85,2.43);line(s,4.31,3.96,4.7,3.12+j*.85,C,1.7,True)
pic(s,V/'raw/current-0.1.12/model-options.png',8.0,2.23,4.36,1.92);pic(s,V/'raw/current-0.1.12/update-panel.png',8.0,4.64,4.36,1.54)
txt(s,'选择模型 / 配置连接 / 查看版本与更新',.9,6.12,6.74,.57,21,I,True);txt(s,'签名校验 → 服务检查 → 自动恢复',.91,6.69,8.78,.33,19,B)
# 30
chapter(6,'用结果，\n检验价值','检查记忆效果，也把积累的资料转化为下一次任务的线索。',['s28'],'metrics')
# 31
s=page('用数据，检验记忆优化','06 / 效果验证',['s28'],'metric-dashboard',True);heading(s,'用数据，检验记忆优化','历史 Debian 兼容环境 · 团队合成数据集',True)
for j,(v,t,b) in enumerate([('50/50','知识检索召回','50组检索全部召回'),('15/15','偏好提取准确','15组偏好全部正确'),('24/25','冲突处理正确','25组中正确处理24组')]):
 x=.83+j*4.18;txt(s,v,x,2.88,3.71,.98,55,F,True);txt(s,t,x+.04,4.13,3.58,.54,25,'C6E6FA',True);txt(s,b,x+.04,4.86,3.58,.46,19,'BED6ED')
rect(s,.86,5.7,3.34,.89,'244F7F');txt(s,'检索 P95  115ms',1.09,5.86,2.95,.42,23,F,True);txt(s,'1000次检索',1.1,6.32,2.9,.24,11,'C6E6FA')
txt(s,'当前麒麟原生版本：\n已完成安装、升级与三端协作验证。',4.68,5.7,7.66,.94,23,F,True)
s.notes_slide.notes_text_frame.text+='\n画面补充：P95 115ms为影片指标画面值，历史原始值115.151ms、1000次检索；历史数据与当前原生功能验证分开。'
# 32
s=page('从记住，到主动提供线索','06 / 洞察与简报',['s26'],'insight-editorial')
txt(s,'从记住，\n到主动提供线索',.74,1.1,6.54,1.61,36,I,True)
film(s,'s26',6.71,1.3,5.69,3.22)
pic(s,'brief',.82,3.36,5.52,3.19)
txt(s,'近期知识',7.09,4.94,4.86,.53,27,B,True);txt(s,'查看线索，回到对应记录。',7.13,5.62,5.12,.46,21,M);txt(s,'按日期回顾采集汇总。',7.13,6.25,5.12,.46,21,M)
# 33
s=page('把个人经验，变成可信协作能力','06 / 应用价值',['s29'],'value-triptych');heading(s,'把个人经验，变成可信协作能力')
for j,(t,b,a) in enumerate([('一次任务','结果留下来','agent-tools-completed'),('一份资料','回答带着来源','directory-source'),('多台电脑','经验接续使用','shared-workspace')]):
 x=.88+j*4.17;txt(s,f'0{j+1}',x,2.23,1.16,.71,43,'82A9D4',True);txt(s,t,x,3.16,3.67,.57,29,B,True);pic(s,a,x,4.13,3.66,1.93);txt(s,b,x,6.38,3.65,.45,23,I,True)
 if j<2:line(s,x+3.69,3.45,x+4.02,3.45,C,2,True)
# 34
s=page('让每一台设备的记忆，彼此相通','结束语',['s30'],'brand-close',True)
for x,y,d in [(1,2.4,.17),(11.8,2.9,.21),(2.3,5.7,.13),(10.3,6.1,.17)]:circle(s,x,y,d,C)
line(s,1.19,2.49,11.77,3.0,'467BA7',1);line(s,2.45,5.77,10.26,6.18,'467BA7',1)
txt(s,'让每一台设备的记忆，\n彼此相通。',1.1,3.01,11.15,1.58,44,F,True,PP_ALIGN.CENTER);txt(s,'PIXIU · 貔貅',4.85,5.46,3.9,.56,27,'C7E3F7',True,PP_ALIGN.CENTER)
# TOC destinations are actual slides; body navigation returns to contents.
for a,p in toc_links:a.click_action.target_slide=prs.slides[p-1]
for n,s in enumerate(list(prs.slides)[2:],3):
 a=txt(s,'目录',11.05,7.1,.61,.27,10,'B6CEE8' if n in [3,8,10,18,22,23,26,30,31,34] else M);a.click_action.target_slide=prs.slides[1]
output=W/'render/项目报告.pptx';prs.save(output)
manifest={'schema':3,'version':'0.1.12','content_authority':'正式宣传片演示增强版','video_sha256':'8e870f0f53af841c67926ae3120cb101564aca438c23d81214c65ffee228a89f','slides':records,'inputs':[{'path':p.relative_to(R).as_posix(),'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in sorted(inputs)],'output':{'path':output.relative_to(R).as_posix(),'sha256':hashlib.sha256(output.read_bytes()).hexdigest()}}
(W/'review/build-manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n');print('Built',len(records),'slides aligned to',len(shots),'video shots')

"""Check the reference trial package, input provenance and rendered-page coverage."""
from pathlib import Path
import hashlib
import json
import re
import subprocess
import zipfile
import xml.etree.ElementTree as ET

from PIL import Image, ImageDraw
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[3]
WORK = ROOT / 'submission/presentation-production'
OUT = WORK / 'render/abc-trial'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


m = json.loads((WORK / 'review/abc-trial-manifest.json').read_text())
ppt = OUT / 'PIXIU项目报告-科技风试作版.pptx'
assert digest(ppt) == m['output_sha256']
for item in m['inputs']:
    assert digest(ROOT / item['path']) == item['sha256'], item['path']
assert {n for s in m['slides'] for n in s['source_pages']} == set(range(1, 32))
with zipfile.ZipFile(ppt) as z:
    assert z.testzip() is None
    xml_parts = []
    for name in z.namelist():
        assert not name.startswith('ppt/embeddings/'), name
        if name.endswith(('.xml', '.rels')):
            data = z.read(name)
            root = ET.fromstring(data)
            xml_parts.append(data.decode())
            if name.endswith('.rels'):
                assert all(e.get('TargetMode') != 'External' for e in root), name
    text = '\n'.join(xml_parts)
    for forbidden in ['ABC', 'XYZ', '人从众', '1234567890', '华南理工',
                      '秦基赫', 'InnoSync', '/home/pluto', 'localhost', '127.0.0.1']:
        assert forbidden not in text, forbidden
    media = {hashlib.sha256(z.read(n)).hexdigest() for n in z.namelist()
             if n.startswith('ppt/media/')}
    for slide in m['slides']:
        for path in slide['screenshots'] + slide.get('illustrations', []):
            assert digest(ROOT / path) in media, path

prs = Presentation(ppt)
assert len(prs.slides) == len(m['slides']) == 32
for slide in m['slides']:
    assert not any(c in slide['title'] for c in '，。？！：'), slide['title']
# Main topics must be present above the body; chapter covers have a separate hierarchy.
for n, (page, entry) in enumerate(zip(prs.slides, m['slides']), 1):
    visible = [(a.text, a.top / 914400) for a in page.shapes if a.has_text_frame]
    matches = [y for t, y in visible if t == entry['title']]
    assert matches, (n, entry['title'])
    assert min(matches) < (3.2 if n in {12, 22} else 1.6), (n, matches)
    assert any(t == f'{n:02}' and 7.1 < y < 7.3 for t,y in visible), (n, 'folio')
    if n in {1,2,12,22,32}:
        assert 'navigation' not in entry
    else:
        nav=entry['navigation']
        header={t for t,y in visible if y < .8}
        assert nav['chapter'] in header and set(nav['tabs']) <= header, (n, nav)
        assert nav['active'] in nav['tabs']
    for t, y in visible:
        assert not any(x in t for x in ['风格试作版', '真实界面：', 'PRODUCT  /  MEMORY']), (n, t)
        assert not (y > 7.1 and ('SDK' in t or '实拍' in t)), (n, t)
# The two product capabilities must remain prominent and consistently named.
for n in [1,2,5,7,32]:
    wording='\n'.join(a.text for a in prs.slides[n-1].shapes if a.has_text_frame)
    assert '记忆焕新' in wording and '记忆互联' in wording, (n, 'flagship names')
assert m['slides'][7]['title'] == '记忆焕新'
assert m['slides'][10]['title'] == '记忆互联'
for page in prs.slides:
    assert all('目录整理' not in a.text for a in page.shapes if a.has_text_frame)
for path, expected in [(OUT / 'PIXIU项目报告-科技风试作版.pdf', 32),
                       (WORK / 'render/reference/ABC公司产品宣传路演PPT.pdf', 20)]:
    info = subprocess.run(['pdfinfo', str(path)], check=True, capture_output=True, text=True).stdout
    assert int(re.search(r'Pages:\s+(\d+)', info).group(1)) == expected, path
rendered = []
for folder, count in [('reference', 20), ('abc-trial', 32)]:
    paths = sorted((WORK / 'render' / folder).glob('slide-*.png'))
    assert [p.name for p in paths] == [f'slide-{i:02}.png' for i in range(1, count + 1)]
    for path in paths:
        with Image.open(path) as im:
            assert im.size == (1600, 900), path
            im.verify()
        rendered.append({'path': str(path.relative_to(ROOT)), 'sha256': digest(path)})

# Preserve the accepted 31-page candidate and its formal asset byte-for-byte.
baseline = '4bf08a7c3b504341cf4082b9f98f63fd2db98a521356674c1eecb36d530eef8e'
assert digest(WORK / 'render/项目报告.pptx') == baseline
assert digest(ROOT / 'docs/delivery/assets/项目报告.pptx') == baseline

# A compact overview complements the full-page images and PDF.
overview = Image.new('RGB', (1600, 8 * 250), '#041F3B')
draw = ImageDraw.Draw(overview)
for i in range(32):
    with Image.open(OUT / f'slide-{i+1:02}.png') as im:
        im = im.resize((400, 225), Image.Resampling.LANCZOS)
        x, y = (i % 4) * 400, (i // 4) * 250
        overview.paste(im, (x, y))
        draw.text((x + 10, y + 230), f'{i+1:02}', fill='#57F1FF')
overview.save(OUT / 'overview.png')
result = {
    'pptx_sha256': digest(ppt),
    'pdf_sha256': digest(OUT / 'PIXIU项目报告-科技风试作版.pdf'),
    'reference_pdf_sha256': digest(WORK / 'render/reference/ABC公司产品宣传路演PPT.pdf'),
    'pages': 32, 'reference_pages': 20,
    'checks': ['ZIP CRC and XML parse', 'all input hashes',
               'all 31 original pages mapped', 'original screenshot bytes embedded',
               'no template identity in XML', 'no embedded workbooks or external links',
               '32 main topics present above body', '32 folios and 27 chapter navigation bars', 'production footnotes removed', 'two flagship names and dedicated feature pages',
               '52 full-page PNGs verified', 'formal candidate and asset unchanged'],
    'visual_review': 'See abc-style-review.md. XML checks cannot inspect raster identity or layout.',
    'rendered_pages': rendered,
}
(WORK / 'review/abc-trial-validation.json').write_text(
    json.dumps(result, ensure_ascii=False, indent=2) + '\n')
print(json.dumps({k: v for k, v in result.items() if k != 'rendered_pages'}, ensure_ascii=False))

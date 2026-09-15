from pathlib import Path
from copy import deepcopy
from io import BytesIO
import zipfile, hashlib, json, posixpath
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
ROOT=Path(__file__).resolve().parents[3]; WORK=ROOT/'submission/presentation-production'
SOURCE=WORK/'source/user-submission-20260915-36pages.pptx'
import argparse
parser=argparse.ArgumentParser(description='Rebuild the reviewed 36-page submission revision into a separate candidate.')
parser.add_argument('--output',type=Path,required=True)
OUT=parser.parse_args().output
assert OUT.resolve()!=SOURCE.resolve(), 'Never overwrite the archived user source'
assert not OUT.exists(), 'Choose a new candidate path; existing files are protected'
OUT.parent.mkdir(parents=True,exist_ok=True)
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
assert SOURCE.is_file(), 'Restore the archived user source before rebuilding'
assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()=='03883a5543085a66d9a6e2afe1943b32e848cbec636770a25f3fcf22932819fa'
r=Presentation(SOURCE)
with zipfile.ZipFile(SOURCE) as z: blobs={n:z.read(n) for n in z.namelist()}
rels=E.fromstring(blobs['ppt/_rels/presentation.xml.rels']); targets={x.get('Id'):x.get('Target') for x in rels}
pr=E.fromstring(blobs['ppt/presentation.xml'])
parts={i:posixpath.normpath('ppt/'+targets[x.get('{'+NS['r']+'}id')]) for i,x in enumerate(pr.find('p:sldIdLst',NS),1)}
def text(s,txt,size=None,accent=()):
 tf=s.text_frame; p0=tf.paragraphs[0]; base=deepcopy(p0.runs[0]._r.rPr) if p0.runs else None
 tf.clear()
 for j,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph()
  run=p.add_run();run.text=line
  if base is not None: run._r.insert(0,deepcopy(base))
  if size:run.font.size=Pt(size)
  if accent: run.font.color.rgb=RGBColor.from_string('FFFFFF')
 # Add intentional emphasis without changing shape geometry.
 if accent:
  for p in tf.paragraphs:
   value=p.text
   for rr in list(p._p.findall('a:r',NS)):p._p.remove(rr)
   import re
   for token in re.split('('+'|'.join(map(re.escape,accent))+')',value):
    if not token:continue
    rr=p.add_run();rr.text=token
    if base is not None:rr._r.insert(0,deepcopy(base))
    if size:rr.font.size=Pt(size)
    rr.font.color.rgb=RGBColor.from_string('FFD29C' if token in accent else 'FFFFFF')
 s.text_frame.word_wrap=True
s=r.slides[5]
updates={8:('国家政策为智能体发展“保驾护航”',20),27:('2025.08',None),28:('2023.07',None),29:('2023.02',None),32:('安全与治理保障',None),35:('《国务院关于深入实施\n“人工智能+”行动的意见》\n国发〔2025〕11号',10),36:('《生成式人工智能服务\n管理暂行办法》\n国家网信办等七部门\n令第15号',10),37:('《数字中国建设整体\n布局规划》\n中共中央、国务院印发',10),42:('将人工智能列为前沿科技攻关重点，为 OS Agent 的记忆优化与智能应用提供战略支撑。',12),43:('推动智能终端与智能体广泛应用，为 OS Agent 的个性化服务与记忆复用拓展应用空间。',12),44:('强调输入与使用记录保护，为 Agent 记忆的授权采集、更正与遗忘提供治理参照。',12),45:('统筹数字技术创新与数据安全，为国产系统上的可信记忆共享与应用融合提供方向。',12)}
for i,(t,sz) in updates.items(): text(s.shapes[i],t,sz,('人工智能','记忆优化','智能体','记忆复用','记录保护','授权采集','更正与遗忘','技术创新','数据安全','可信记忆共享') if i>=42 else ())
# Use the existing footer shape; policy page shape count and all geometry remain unchanged.
f=s.shapes[4].shapes[1];text(f,'06',10)
f.text_frame.margin_top=f.text_frame.margin_bottom=0
f.text_frame.paragraphs[0].alignment=PP_ALIGN.CENTER
for rr in f.text_frame.paragraphs[0].runs:rr.font.color.rgb=RGBColor.from_string('BCD8EA')
policy_rel=posixpath.dirname(parts[6])+'/_rels/'+posixpath.basename(parts[6])+'.rels'
rr=E.fromstring(blobs[policy_rel])
for idx,num in [(11,'01'),(12,'02'),(13,'03')]:
 im=next((WORK/'reference/赛题政策素材/截图').glob(num+'*'))
 name='pixiu-policy-'+num+'.png'; blobs['ppt/media/'+name]=im.read_bytes()
 rid='rIdPixiuPolicy'+num
 E.SubElement(rr,'{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',Id=rid,Type=NS['r']+'/image',Target='../media/'+name)
 shape=s.shapes[idx];shape._element.find('.//a:blip',NS).set('{'+NS['r']+'}embed',rid)
 # Use native picture crop. The full screenshot stays in the same existing image frame.
 for key in ['crop_left','crop_right','crop_top','crop_bottom']:setattr(shape,key,0)
blobs[policy_rel]=E.tostring(rr,xml_declaration=True,encoding='UTF-8',standalone=True)
CYAN='4BE5F4'; GOLD='FFD29C'; WHITE='FFFFFF'; NAVY='073A68'
def box(sl,x,y,w,h,txt='',size=14,color=WHITE,fill=None,line=None,bold=False):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h)) if fill else sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
 if fill:
  sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill)
  sh.line.color.rgb=RGBColor.from_string(line or fill);sh.line.width=Pt(1)
  sh.adjustments[0]=.12
 tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.08 if fill else 0);tf.margin_top=tf.margin_bottom=0
 tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 for j,st in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.alignment=PP_ALIGN.CENTER if fill else PP_ALIGN.LEFT
  p.space_before=p.space_after=0
  q=p.add_run();q.text=st;q.font.name='Microsoft YaHei';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=RGBColor.from_string(color)
 return sh
def line(sl,x1,y1,x2,y2,arrows=False,dash=False):
 sh=sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
 sh.line.color.rgb=RGBColor.from_string(CYAN);sh.line.width=Pt(2)
 ln=sh._element.find('.//a:ln',NS)
 if dash:E.SubElement(ln,'{'+NS['a']+'}prstDash',val='dash')
 if arrows:
  E.SubElement(ln,'{'+NS['a']+'}headEnd',type='triangle',w='sm',len='sm');E.SubElement(ln,'{'+NS['a']+'}tailEnd',type='triangle',w='sm',len='sm')
 return sh
# Architecture promotion: retain each independent screenshot and its editable annotation.
s=r.slides[11]
for base,x in [(20,.55),(27,5.0),(34,9.45)]:
 card=s.shapes[base];card.left=Inches(x);card.top=Inches(2.8);card.width=Inches(3.33);card.height=Inches(3.53)
 for idx,y,h in [(base+1,2.97,.4),(base+2,3.40,.3),(base+6,6.06,.23)]:
  sh=s.shapes[idx];sh.left=Inches(x+.10);sh.top=Inches(y);sh.width=Inches(3.13);sh.height=Inches(h)
 # Scale screenshot and highlight together so annotation keeps its meaning.
 im=s.shapes[base+4];oldx,oldy=im.left,im.top;factor=3.13/(im.width/914400)
 for idx in [base+4,base+5]:
  sh=s.shapes[idx];sh.left=Inches(x+.10)+round((sh.left-oldx)*factor);sh.top=Inches(4.23)+round((sh.top-oldy)*factor);sh.width=round(sh.width*factor);sh.height=round(sh.height*factor)
 strip=s.shapes[base+3];strip.left=Inches(x+.10);strip.top=Inches(4.15);strip.width=Inches(3.13);strip.height=Inches(.14)
 box(s,x+.18,3.82,2.97,.30,'Agent · PIXIU 本地记忆',12,CYAN,NAVY,CYAN,True)
 line(s,x+1.665,4.12,x+1.665,4.22)
for x1,x2,label in [(3.88,5.0,'对等同步'),(8.33,9.45,'对等同步')]:
 line(s,x1,3.98,x2,3.98,True)
 box(s,x1,3.62,x2-x1,.26,label,11,CYAN)
# Lower bypass makes the topology a peer network, rather than a mandatory central relay.
line(s,2.215,6.33,2.215,6.64,False,True)
line(s,2.215,6.64,11.115,6.64,True,True)
line(s,11.115,6.33,11.115,6.64,False,True)
box(s,4.50,6.47,4.32,.34,'可信设备按范围共享 · 离线重连补齐',12,CYAN,NAVY,None,True)
s.shapes[41].top=Inches(6.90);s.shapes[41].height=Inches(.31)
for p in s.shapes[41].text_frame.paragraphs:
 for q in p.runs:q.font.size=Pt(20)
s.shapes[42].top=Inches(7.32)
# Explanatory copy in the previously empty lower-left region; keep bibliography untouched.
s=r.slides[24]
box(s,.48,4.06,5.95,.38,'研究启发与 PIXIU 的工程落点',21,CYAN,bold=True)
rows=[('01  结构化记忆','关联事实、实体与来源，让召回结果可追溯、可核对。'),('02  持续整合','随新资料提出更正与合并，经用户审批后更新记忆。'),('03  共享与上下文','按授权范围同步副本，为后续任务选取相关记忆。')]
for j,(title,body) in enumerate(rows):
 y=4.65+j*.67
 box(s,.48,y,.06,.52,fill=CYAN)
 box(s,.68,y,5.65,.27,title,16,GOLD,bold=True)
 box(s,.68,y+.30,5.65,.26,body,12.5,WHITE)
box(s,.68,6.75,5.65,.26,'以记忆组织、更新与复用为主线，衔接国产端侧应用。',12,CYAN)
# Existing visible footers use physical slide numbers; covers remain unnumbered.
for n,sl in enumerate(r.slides,1):
 for sh in sl.shapes:
  if sh.has_text_frame and sh.top>Inches(7.0) and sh.left>Inches(11.5) and sh.text.strip().isdigit():
   if sh.text.strip()!=f'{n:02}':text(sh,f'{n:02}')
# Update only slide parts with actual changes. Untouched media, charts and masters retain exact bytes.
changed=[]
for n,sl in enumerate(r.slides,1):
 old=E.fromstring(blobs[parts[n]])
 if E.tostring(old)!=E.tostring(sl._element):
  blobs[parts[n]]=E.tostring(sl._element,xml_declaration=True,encoding='UTF-8',standalone=True);changed.append(n)
with zipfile.ZipFile(OUT,'w',zipfile.ZIP_DEFLATED) as z:
 for name,b in blobs.items():z.writestr(name,b)
print('updated',OUT,'slides',changed)

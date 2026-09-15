"""Draw editable logic diagrams inside three specified regions of the latest deck."""
from pathlib import Path
import argparse,zipfile
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.shapes import MSO_SHAPE as K, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.dml.color import RGBColor
P=argparse.ArgumentParser();P.add_argument('--source',type=Path,required=True);P.add_argument('--output',type=Path,required=True);a=P.parse_args();assert not a.output.exists()
r=Presentation(a.source);NS={'a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main'}
C='4CE6F5';G='FFD298';W='E9F8FF';D='062C50';F='08507A'
def shape(s,x,y,w,h,kind=K.ROUNDED_RECTANGLE,fill=D,line=C):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(1)
 else:q.line.fill.background()
 if kind==K.ROUNDED_RECTANGLE:q.adjustments[0]=.1
 return q
def text(s,x,y,w,h,t,size=13,color=W,bold=False,align=PP_ALIGN.CENTER):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=q.text_frame;tf.clear();tf.word_wrap=True;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.vertical_anchor=MSO_ANCHOR.MIDDLE
 for j,v in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if j==0 else tf.add_paragraph();p.alignment=align;p.space_before=p.space_after=0
  rr=p.add_run();rr.text=v;rr.font.name='Microsoft YaHei';rr.font.size=Pt(size);rr.font.bold=bold;rr.font.color.rgb=RGBColor.from_string(color)
 return q
def node(s,x,y,w,h,t,size=12,kind=K.ROUNDED_RECTANGLE,color=C):
 q=shape(s,x,y,w,h,kind,line=color);text(s,x+.025,y+.02,w-.05,h-.04,t,size,color,True);return q
def edge(s,points,arrow=True,color=C,dash=False):
 for j,((x,y),(u,v)) in enumerate(zip(points,points[1:])):
  q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(u),Inches(v));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(1.6)
  ln=q._element.find('.//a:ln',NS)
  if dash:E.SubElement(ln,'{'+NS['a']+'}prstDash',val='dash')
  if arrow and j==len(points)-2:E.SubElement(ln,'{'+NS['a']+'}tailEnd',type='triangle',w='sm',len='sm')
def remove(s,indices):
 orig=list(s.shapes)
 for i in indices:s.shapes._spTree.remove(orig[i]._element)
def doc(s,x,y,w,h,color=C):
 shape(s,x,y,w,h,K.FOLDED_CORNER,D,color)
 for j in range(3):edge(s,[(x+w*.18,y+h*(.36+j*.17)),(x+w*.76,y+h*(.36+j*.17))],False,color)
def database(s,x,y,w,h):
 shape(s,x,y,w,h,K.CAN,D,C)
 edge(s,[(x+w*.18,y+h*.56),(x+w*.82,y+h*.56)],False)
def chat(s,x,y,w,h):
 shape(s,x,y,w,h,K.ROUNDED_RECTANGULAR_CALLOUT,D,C)
 edge(s,[(x+w*.18,y+h*.35),(x+w*.81,y+h*.35)],False)
 edge(s,[(x+w*.18,y+h*.60),(x+w*.62,y+h*.60)],False)
# 19: state flow with explicit transition causes and a retrieval return path.
s=r.slides[18];assert s.shapes[2].text=='记忆阶段流转';remove(s,range(26,35))
for y,title,body,icon in [(2.30,'短期 · 当前任务','本轮问题与处理要点',chat),(3.95,'中期 · 阶段状态','阶段内容与到期时间',doc),(5.60,'长期 · 持久知识','知识结构与检索索引',database)]:
 shape(s,8.42,y,3.84,.78,fill=F);icon(s,8.59,y+.16,.50,.43)
 text(s,9.25,y+.08,2.82,.30,title,17,G,True,PP_ALIGN.LEFT)
 text(s,9.25,y+.43,2.82,.25,body,12,W,False,PP_ALIGN.LEFT)
edge(s,[(10.32,3.08),(10.32,3.95)])
text(s,10.55,3.24,1.82,.42,'压缩／切换\n会话结束',11,C)
edge(s,[(10.32,4.73),(10.32,5.60)])
text(s,10.55,4.98,1.82,.32,'选择长期保留',11,G)
edge(s,[(8.42,6.0),(7.65,6.0),(7.65,2.69),(8.42,2.69)],True,C,True)
shape(s,7.14,3.85,1.02,.75,fill='06304F',line=None)
text(s,7.17,3.89,.96,.65,'检索召回\n供任务复用',11,C,True)
# 20: retain panel outlines/titles, replace the two paragraphs with branching diagrams.
s=r.slides[19];assert s.shapes[2].text=='双层冲突治理';remove(s,[26,30])
# Same record has two versions; version comparison handles causal vs concurrent updates.
node(s,1.00,2.68,.78,.26,'版本 A',10)
node(s,1.00,3.08,.78,.26,'版本 B',10)
edge(s,[(1.78,2.81),(1.96,2.81),(1.96,3.00),(2.18,3.00)])
edge(s,[(1.78,3.21),(1.96,3.21),(1.96,3.00)],False)
node(s,2.18,2.64,1.05,.69,'版本向量\n判断因果',10.5,K.DIAMOND)
edge(s,[(3.23,2.98),(3.44,2.98),(3.44,2.79),(3.78,2.79)])
edge(s,[(3.44,2.98),(3.44,3.19),(3.78,3.19)])
node(s,3.78,2.63,1.16,.30,'因果更新',11)
node(s,3.78,3.04,1.16,.30,'并发 LWW',10.5,color=G)
edge(s,[(4.94,2.79),(5.10,2.79),(5.10,2.99),(5.32,2.99)])
edge(s,[(4.94,3.19),(5.10,3.19),(5.10,2.99)],False)
node(s,5.32,2.69,1.00,.59,'一致\n副本',12,K.CAN)
# Semantic conflicts choose between three outcomes, rather than a linear pipeline.
doc(s,7.02,2.70,.30,.39);doc(s,7.16,2.80,.30,.39)
text(s,7.53,2.74,1.00,.40,'新旧知识',11,W,True)
edge(s,[(8.55,2.96),(8.82,2.96)])
node(s,8.82,2.62,1.19,.70,'实体／字段\n比较',10.5,K.DIAMOND)
edge(s,[(10.01,2.96),(10.30,2.96)],False)
for y,t in [(2.60,'更新'),(2.86,'合并'),(3.12,'人工确认')]:
 edge(s,[(10.30,2.96),(10.30,y+.115),(10.64,y+.115)])
 node(s,10.64,y,1.51,.23,t,10.5,color=G if t=='人工确认' else C)
# 21: online fan-out with ACK feedback; anti-entropy compares missing operations.
s=r.slides[20];assert s.shapes[2].text=='去中心化同步协议';remove(s,range(24,30))
shape(s,.8,2.76,5.52,1.78,fill='064671',line=None)
text(s,1.0,2.87,4.90,.31,'在线 · 更新传播',20,G,True,PP_ALIGN.LEFT)
doc(s,1.16,3.36,.39,.51)
text(s,.96,3.91,.87,.26,'签名操作',11,W,True)
edge(s,[(1.66,3.62),(3.02,3.62),(3.02,3.44),(4.31,3.44)])
edge(s,[(3.02,3.62),(3.02,4.05),(4.31,4.05)])
shape(s,2.28,3.49,.12,.12,K.OVAL,C,None)
text(s,2.0,3.20,1.75,.27,'Gossip 推送',12,C,True)
for y,name in [(3.24,'对端 B'),(3.85,'对端 C')]:
 node(s,4.31,y,1.73,.43,name+' · 物化',11)
edge(s,[(4.31,4.26),(3.82,4.26),(3.82,4.40),(1.34,4.40),(1.34,4.22)],True,G,True)
text(s,2.10,4.13,1.35,.23,'ACK 接收确认',10,G)
shape(s,.8,4.75,5.52,1.74,fill='064671',line=None)
text(s,1.0,4.87,4.90,.31,'重连 · 反熵对账',20,G,True,PP_ALIGN.LEFT)
text(s,1.00,5.30,1.03,.25,'本机摘要',11,W,True)
text(s,1.00,5.82,1.03,.25,'对端摘要',11,W,True)
for row,y in enumerate([5.30,5.82]):
 for j in range(4):
  missing=row==1 and j==2
  node(s,2.17+j*.40,y,.31,.27,'—' if missing else str(j+1),10,color=G if j==2 else C)
edge(s,[(3.00,5.59),(3.00,5.79)],True,G)
text(s,3.93,5.28,1.93,.25,'识别缺失操作 3',11,G,True)
edge(s,[(3.80,5.95),(4.13,5.95)])
node(s,4.13,5.75,1.85,.43,'补齐 · CRDT 合并',11)
text(s,1.20,6.22,4.80,.19,'交换摘要  /  按差异补齐  /  合并副本',10,C)
with zipfile.ZipFile(a.source) as z:blobs={n:z.read(n) for n in z.namelist()}
for n in [19,20,21]:blobs[str(r.slides[n-1].part.partname).lstrip('/')]=E.tostring(r.slides[n-1]._element,xml_declaration=True,encoding='UTF-8',standalone=True)
with zipfile.ZipFile(a.output,'w',zipfile.ZIP_DEFLATED) as z:
 for n,b in blobs.items():z.writestr(n,b)
print(a.output)

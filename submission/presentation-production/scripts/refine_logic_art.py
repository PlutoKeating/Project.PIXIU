"""Refine the three logic diagrams with editable, layered native PPT geometry.

Input is the 36-page deck after logic_diagrams.py; only slide 19–21 XML changes.
Always generate a separate candidate and merge against the latest saved deck.
"""
import argparse, zipfile
from pathlib import Path
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE as K, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

A='http://schemas.openxmlformats.org/drawingml/2006/main'
C='56E8F5'; G='FFD29A'; W='EAF9FF'; M='9EC4D8'
def sub(p,t,**kw):return E.SubElement(p,'{'+A+'}'+t,**kw)
def gradient(q,colors=('116987','063562'),angle=0):
 p=q._element.spPr
 for el in list(p):
  if E.QName(el).localname in ['solidFill','noFill','gradFill']:p.remove(el)
 g=E.Element('{'+A+'}gradFill',rotWithShape='1');gs=sub(g,'gsLst')
 for pos,col in [(0,colors[0]),(100000,colors[1])]:sub(sub(gs,'gs',pos=str(pos)),'srgbClr',val=col)
 sub(g,'lin',ang=str(angle*60000),scaled='1');p.insert(2,g)
def shadow(q):
 e=sub(q._element.spPr,'effectLst');o=sub(e,'outerShdw',blurRad='85000',dist='45000',dir='5400000',algn='ctr',rotWithShape='0');sub(sub(o,'srgbClr',val='00152B'),'alpha',val='50000')
def shape(s,x,y,w,h,kind=K.ROUNDED_RECTANGLE,fill='08375D',line=None,grad=None,sh=False):
 q=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h));q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(.7)
 else:q.line.fill.background()
 if kind==K.ROUNDED_RECTANGLE:q.adjustments[0]=.13
 if grad:gradient(q,grad)
 if sh:shadow(q)
 return q
def text(s,x,y,w,h,t,size=13,color=W,bold=False,align=PP_ALIGN.LEFT):
 q=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));f=q.text_frame;f.clear();f.word_wrap=True;f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0;f.vertical_anchor=MSO_ANCHOR.MIDDLE
 for j,v in enumerate(t.split('\n')):
  p=f.paragraphs[0] if j==0 else f.add_paragraph();p.alignment=align;p.space_before=p.space_after=0
  r=p.add_run();r.text=v;r.font.name='Microsoft YaHei';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor.from_string(color)
 return q
def edge(s,pts,color=C,width=1.3,arrow=True,dash=False):
 for j,((x,y),(u,v)) in enumerate(zip(pts,pts[1:])):
  q=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x),Inches(y),Inches(u),Inches(v));q.line.color.rgb=RGBColor.from_string(color);q.line.width=Pt(width)
  ln=q._element.spPr.find('{'+A+'}ln')
  if dash:sub(ln,'prstDash',val='dash')
  if arrow and j==len(pts)-2:sub(ln,'tailEnd',type='triangle',w='sm',len='sm')
def poly(s,pts,fill,line=None):
 b=s.shapes.build_freeform(Inches(pts[0][0]),Inches(pts[0][1]));b.add_line_segments([(Inches(x),Inches(y)) for x,y in pts[1:]],close=True);q=b.convert_to_shape();q.fill.solid();q.fill.fore_color.rgb=RGBColor.from_string(fill)
 if line:q.line.color.rgb=RGBColor.from_string(line);q.line.width=Pt(.6)
 else:q.line.fill.background()
 return q
def platform(s,x,y,w,h=.20):
 poly(s,[(x,y),(x+w-.18,y),(x+w,y+h),(x+.18,y+h)],'0A4967','267491')
 poly(s,[(x+.18,y+h),(x+w,y+h),(x+w,y+h+.07),(x+.18,y+h+.07)],'062F50')
 edge(s,[(x+.30,y+h),(x+w-.23,y+h)],C,.7,False)
def chip(s,x,y,w,h,t,size=11,color=C):
 shape(s,x,y+.025,w,h,fill='031E3D')
 shape(s,x,y,w,h,line='32809A' if color==C else 'AD9067',grad=('125876','072E53') if color==C else ('675C49','18374B'),sh=True)
 text(s,x+.04,y+.01,w-.08,h-.02,t,size,W,True,PP_ALIGN.CENTER)
def icon(s,x,y,d,kind,color=C):
 shape(s,x,y+.05,d,d,K.OVAL,fill='031F3A')
 shape(s,x,y,d,d,K.OVAL,line='2C748E',grad=('197995','073455'),sh=True)
 shape(s,x+.07,y+.07,d-.14,d-.14,K.OVAL,line=color,fill='093B59')
 a=x+d*.28;b=y+d*.26;w=d*.44;h=d*.46
 if kind=='db':
  shape(s,a,b,w,h,K.CAN,grad=('80DCE2','16678B'),line=color)
  edge(s,[(a+.035,b+h*.65),(a+w-.035,b+h*.65)],color,.65,False)
 elif kind=='chat':
  shape(s,a,b,w,h,K.ROUNDED_RECTANGULAR_CALLOUT,grad=('70D6DF','176087'),line=color)
  edge(s,[(a+w*.2,b+h*.35),(a+w*.8,b+h*.35)],'D7FCFF',.8,False)
  edge(s,[(a+w*.2,b+h*.6),(a+w*.6,b+h*.6)],'D7FCFF',.8,False)
 else:
  shape(s,a+.045,b-.045,w,h,K.RECTANGLE,fill='145471',line='33879B')
  shape(s,a,b,w,h,K.FOLDED_CORNER,grad=('A1E6EB','277A98'),line=color)
  for k in range(2):edge(s,[(a+w*.18,b+h*(.40+k*.22)),(a+w*.75,b+h*(.40+k*.22))],'E3FFFF',.7,False)
def panel(s,x,y,w,h,title,num):
 shape(s,x,y,w,h,line='1E6685',grad=('0B486C','06274D'),sh=True)
 edge(s,[(x+.22,y+.01),(x+1.17,y+.01)],C,2,False)
 text(s,x+.23,y+.11,w-1,.34,title,19,G,True)
 text(s,x+w-.67,y+.10,.45,.34,num,15,'4287A6',True,PP_ALIGN.RIGHT)
def terminal(s,x,y,label):
 platform(s,x-.11,y+.40,.76,.13)
 shape(s,x,y,.56,.40,grad=('297E9A','082F56'),line='5FC7DC',sh=True)
 shape(s,x+.045,y+.045,.47,.27,grad=('126285','052340'))
 for j in range(2):edge(s,[(x+.11,y+.12+j*.075),(x+.38,y+.12+j*.075)],C,.9,False)
 edge(s,[(x+.28,y+.40),(x+.28,y+.49)],C,2,False)
 text(s,x+.68,y+.03,1.01,.35,label,11.5,W,True)
 text(s,x+.68,y+.33,1.01,.22,'本地物化',9.5,M)

def build(r):
 # Keep all user-owned elements; remove only shapes generated in the previous pass.
 for n,threshold in [(19,94),(20,91),(21,90)]:
  s=r.slides[n-1]
  for q in list(s.shapes):
   if q.shape_id>=threshold:s.shapes._spTree.remove(q._element)
 # Memory: suspended tiers, illuminated timeline, retrieval feedback rail.
 s=r.slides[18]
 edge(s,[(7.42,5.86),(7.42,2.72),(7.94,2.72)],'287A98',1,True)
 edge(s,[(8.65,5.86),(7.42,5.86)],'287A98',1,False)
 shape(s,7.12,4.04,.62,.92,grad=('0E5674','082F51'))
 text(s,7.17,4.12,.52,.76,'检索\n召回',11,C,True,PP_ALIGN.CENTER)
 for k,(x,y,title,body,kind) in enumerate([(8.00,2.34,'短期 · 当前任务','本轮问题与处理要点','chat'),(8.30,3.95,'中期 · 阶段状态','阶段内容与到期时间','doc'),(8.60,5.55,'长期 · 持久知识','知识结构与检索索引','db')]):
  platform(s,x-.04,y+.75,4.0,.20)
  shape(s,x+.30,y+.08,3.63,.69,grad=('135677','07305A'),sh=True)
  edge(s,[(x+.81,y+.09),(x+3.56,y+.09)],'35829A',.6,False)
  icon(s,x,y-.05,.84,kind)
  text(s,x+1.00,y+.11,2.84,.30,title,17,G,True)
  text(s,x+1.00,y+.46,2.78,.24,body,11.6,W)
  text(s,x-.42,y+.01,.29,.24,'0'+str(k+1),10,C,True,PP_ALIGN.CENTER)
  if k<2:
   xx=x+1.91
   edge(s,[(xx,y+1.04),(xx,y+1.45),(xx+.30,y+1.45),(xx+.30,y+1.57)],C,1.4)
   text(s,xx+.23,y+1.03,2.10,.33,'压缩／切换／会话结束' if k==0 else '选择长期保留',10.4,M if k==0 else G)
 # Conflict: two polished decision rails. Original titles/backgrounds are retained.
 s=r.slides[19]
 for q in s.shapes:
  if q.shape_id in [75,79]:gradient(q,('0A496B','062B50'))
 # tiny raised version sheets
 for y,t in [(2.70,'版本 A'),(3.06,'版本 B')]:
  shape(s,1.09,y-.035,.68,.26,fill='1A6480')
  chip(s,1.03,y,.70,.25,t,9.6)
 edge(s,[(1.77,2.83),(1.96,2.83),(1.96,3.00),(2.19,3.00)])
 edge(s,[(1.77,3.19),(1.96,3.19),(1.96,3.00)],arrow=False)
 shape(s,2.18,2.65,1.02,.70,K.HEXAGON,line='59D6E7',grad=('25738C','083151'),sh=True)
 text(s,2.28,2.73,.82,.45,'版本向量\n判断因果',10.6,W,True,PP_ALIGN.CENTER)
 for y,t,col in [(2.70,'因果更新',C),(3.08,'并发 LWW',G)]:
  edge(s,[(3.20,3.0),(3.44,3.0),(3.44,y+.13),(3.72,y+.13)],col,1.0)
  chip(s,3.72,y,1.12,.27,t,10.3,col)
  edge(s,[(4.84,y+.13),(5.02,y+.13),(5.02,3.0),(5.28,3.0)],col,1.0)
 platform(s,5.27,3.23,1.02,.10)
 shape(s,5.37,2.70,.79,.54,K.CAN,grad=('6CBFCD','135274'),line='6EDAE8',sh=True)
 text(s,5.39,2.84,.75,.28,'一致副本',10,'F0FFFF',True,PP_ALIGN.CENTER)
 # Right: layered knowledge stack, central comparison lens, three action routes.
 for dx,dy in [(.10,-.04),(.05,0),(0,.04)]:shape(s,7.00+dx,2.73+dy,.34,.43,K.FOLDED_CORNER,grad=('74C8D7','164D74'),line='5EA9C0')
 text(s,7.53,2.85,.89,.27,'新旧知识',10.6,W,True)
 edge(s,[(8.40,3.00),(8.70,3.00)],width=1)
 shape(s,8.70,2.66,1.18,.68,K.HEXAGON,line='59D6E7',grad=('25738C','083151'),sh=True)
 text(s,8.82,2.76,.94,.42,'实体／字段\n比较',10.5,W,True,PP_ALIGN.CENTER)
 for y,t,col in [(2.61,'更新',C),(2.88,'合并',C),(3.15,'人工确认',G)]:
  edge(s,[(9.88,3.00),(10.13,3.00),(10.13,y+.105),(10.43,y+.105)],col,1)
  shape(s,10.43,y,.22,.21,K.OVAL,grad=('27859A','104568') if col==C else ('B18F5F','645341'))
  text(s,10.44,y,.20,.20,'✓' if t!='人工确认' else '!',8,W,True,PP_ALIGN.CENTER)
  text(s,10.78,y,1.21,.21,t,11,W if col==C else G,True)
 # Sync: three-dimensional devices and a routed fan-out, followed by operation tiles.
 s=r.slides[20]
 panel(s,.8,2.76,5.52,1.79,'在线 · 更新传播','01')
 icon(s,1.09,3.36,.69,'doc')
 text(s,1.01,4.13,.94,.24,'签名操作',11,W,True,PP_ALIGN.CENTER)
 # central routing hub, layered oval pedestal
 shape(s,2.59,3.75,.75,.21,K.OVAL,fill='03213B',line='176A86')
 shape(s,2.61,3.69,.71,.21,K.OVAL,grad=('258CA6','0C3F62'),line='47C2D5')
 shape(s,2.83,3.46,.28,.31,K.HEXAGON,grad=('85ECF0','247F9C'),line=C)
 edge(s,[(1.81,3.70),(2.60,3.70)],width=1.6)
 edge(s,[(3.33,3.69),(3.64,3.69),(3.64,3.44),(4.12,3.44)],width=1.3)
 edge(s,[(3.64,3.69),(3.64,4.08),(4.12,4.08)],width=1.3)
 text(s,2.10,3.21,1.72,.24,'Gossip 推送',11.5,C,True,PP_ALIGN.CENTER)
 terminal(s,4.14,3.24,'对端 B');terminal(s,4.14,3.91,'对端 C')
 edge(s,[(4.0,4.36),(2.14,4.36),(2.14,4.09),(1.83,4.09)],G,.9,True,True)
 text(s,2.47,4.09,1.36,.21,'ACK 接收确认',9.5,G)
 panel(s,.8,4.77,5.52,1.73,'重连 · 反熵对账','02')
 for row,y in enumerate([5.37,5.96]):
  text(s,1.00,y+.02,.90,.25,'本机摘要' if row==0 else '对端摘要',10.5,W,True)
  platform(s,1.95,y+.24,1.86,.10)
  for j in range(4):
   missing=row==1 and j==2;x=2.06+j*.43
   if missing:
    q=shape(s,x,y,.32,.30,fill='092D46',line=G)
    ln=q._element.spPr.find('{'+A+'}ln');sub(ln,'prstDash',val='dash')
   else:shape(s,x,y,.32,.30,grad=('55A8B8','175673') if j!=2 else ('D6B784','70603F'),line='6FBCCA' if j!=2 else G,sh=True)
   text(s,x,y,.32,.29,'−' if missing else str(j+1),10.8,W,True,PP_ALIGN.CENTER)
 edge(s,[(3.08,5.75),(3.08,5.93)],G,1.5)
 text(s,3.41,5.75,.41,.18,'缺失',8.4,G)
 edge(s,[(3.93,6.13),(4.23,6.13)],G,1.2)
 icon(s,4.53,5.40,.64,'db')
 text(s,4.01,6.10,1.98,.26,'补齐 · CRDT 合并',11,W,True,PP_ALIGN.CENTER)
 text(s,4.10,5.13,1.85,.22,'识别差异 · 按需补齐',9.5,G,False,PP_ALIGN.CENTER)

if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args();assert not a.output.exists()
 r=Presentation(a.source);assert len(r.slides)==36
 for n,t in [(19,'记忆阶段流转'),(20,'双层冲突治理'),(21,'去中心化同步协议')]:assert r.slides[n-1].shapes[2].text==t
 build(r)
 with zipfile.ZipFile(a.source) as z:
  entries=z.infolist();blobs={i.filename:z.read(i.filename) for i in entries}
 for n in [19,20,21]:blobs[f'ppt/slides/slide{n}.xml']=E.tostring(r.slides[n-1]._element,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(a.output,'w') as z:
  for i in entries:z.writestr(i,blobs[i.filename])
 print(a.output)

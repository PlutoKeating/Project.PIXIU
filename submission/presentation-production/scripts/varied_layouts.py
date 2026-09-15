"""Content-led layouts using native decorative objects copied from the ABC deck."""
from copy import deepcopy
from io import BytesIO
from pptx import Presentation
from pptx.util import Inches
from lxml import etree as ET
from product_screenshots import Canvas, window, NOTES, SOURCES, CAP, WORK, NS

REF=Presentation(WORK/'reference/ABC公司产品宣传路演PPT.pptx')
PAGES={7,10,11,12,13,15,16,17,18,19,20,21,22,23,25,26,27,28,30}
C='57F1FF';G='FFDCAB'

def art(c,page,index,x,y,w,h):
    src=REF.slides[page-1].shapes[index]
    if src.shape_type==13:
        c.slide.shapes.add_picture(BytesIO(src.image.blob),Inches(x),Inches(y),Inches(w),Inches(h))
    else:
        e=deepcopy(src._element)
        assert not e.xpath('.//a:blip')
        for t in e.xpath('.//a:t'):t.text=''
        for n in e.xpath('.//p:cNvPr'):n.set('id',str(c.slide.shapes._next_shape_id));n.set('name','ABC native decoration')
        xf=e.find('p:spPr/a:xfrm',NS)
        assert xf is not None
        for name,a,b in [('off',x,y),('ext',w,h)]:
            el=xf.find('a:'+name,NS)
            for key,v in zip(('x','y') if name=='off' else ('cx','cy'),(a,b)):el.set(key,str(Inches(v)))
        c.slide.shapes._spTree.insert_element_before(e,'p:extLst')
    c.artworks.append({'reference_page':page,'shape_index':index,'box':[x,y,w,h]})

def floor(c,y=6.86):art(c,11,26,.55,y,12.25,.15)
def glow(c,x,y,w):art(c,3,59,x,y,w,.16)
def panel(c,x,y,w,h,kind='card'):
    p,i={'card':(3,18),'wing':(12,1),'round':(11,1),'band':(7,34)}[kind]
    art(c,p,i,x,y,w,h)
def block(c,title,body,x,y,w,h=1.15,size=17):
    c.text(title,x,y,w,.38,20,G,True)
    c.text(body,x,y+.48,w,h-.48,size)
def card(c,title,body,x,y,w,h=1.5):
    panel(c,x,y,w,h);glow(c,x+.25,y-.04,w-.5)
    c.text(title,x+.2,y+.12,w-.4,.33,17 if h<=1.4 else 20,G,True)
    c.text(body,x+.2,y+.55,w-.4,h-.60,14 if h<=1.4 else 16)
def shot(c,key,x,y,w,h,caption):window(c,key,x,y,w,h,caption)
def intro(c,t):c.text(t,.85,1.95,11.65,.58,17)
def steps(c,items,x,y,w,h):
    gap=.2;cw=(w-gap*(len(items)-1))/len(items)
    for i,(a,b) in enumerate(items):card(c,a,b,x+i*(cw+gap),y,cw,h)

def compose(page):
    c=Canvas();c.artworks=[]
    c.layout=''
    if page==7:
        c.layout='product stage with feature ribbons'
        panel(c,.65,2.13,8.1,4.55,'wing');glow(c,1.3,2.1,6.6)
        shot(c,'bill',.9,2.28,7.6,3.94,'图 1：助手在会话中检索并核对账单')
        for i,(t,b) in enumerate(NOTES[7][:2]):
            panel(c,9.0,2.38+i*1.67,3.5,1.48,'band');block(c,t,b,9.18,2.52+i*1.67,3.16,1.28,16)
        c.text('回答附来源，结果可核对',8.95,6.08,3.55,.44,22,C,True)
        floor(c)
    elif page==10:
        c.layout='two flagship podiums'
        panel(c,.5,2.25,12.35,4.52,'wing')
        for i,(key,t,b) in enumerate([('devices','记忆共享，分布互连','授权共享 · 本地副本 · 重连补齐'),('dreaming','自动记忆，持续整合','后台理解 · 关联记忆 · 更正审批')]):
            x=.85+i*6.05
            c.text('0'+str(i+1),x,2.02,.85,.65,36,C,True)
            c.text(t,x+.94,2.13,4.65,.42,23,G,True)
            shot(c,key,x,2.84,5.6,2.96,'图 '+str(i+1)+'：'+('用户配置可信设备与共享范围' if i==0 else '用户核对后台生成的更正方案'))
            panel(c,x,6.27,5.6,.49,'band');c.text(b,x+.14,6.36,5.3,.32,17,C,True,True)
        floor(c)
    elif page in {11,30}:
        c.layout='three device columns' if page==11 else 'three-device verification stage'
        intro(c,'可信设备按授权范围共享，各端保留本地副本并在重连后补齐差异。' if page==11 else '原验证包含两端并发与第三端离线，下图分别展示三台虚拟机当前的共享记录。')
        if page==30:panel(c,.5,2.64,12.3,3.98,'wing')
        for i,(key,domain) in enumerate([('a','kylin-v11'),('b','pixiu-video-b'),('c','pixiu-release-c')]):
            x=.8+i*4.0
            if page==11:panel(c,x,2.68,3.73,3.52,'card')
            c.text('设备 '+key.upper(),x,2.86,3.73,.42,23,C,True,True)
            c.text(domain,x,3.36,3.73,.25,13,center=True)
            SOURCES[key]=(CAP/('device-'+key+'-memory.png'),(0,0,3840,2160),[(1520,875,1100,100) if key=='c' else (1770,1035,1150,105)],'共享记录正文')
            glow(c,x+.1,3.75,3.53)
            shot(c,key,x+.05,3.83,3.63,2.05,'图 '+str(i+1)+'：设备 '+key.upper()+' 查询同一条共享记忆')
        c.text('三端一致：11 月 10 日下午三点 · 三楼档案室',.9,6.46,11.5,.42,23,G,True,True);floor(c)
    elif page==12:
        c.layout='vertical integration process beside full review'
        intro(c,'授权资料变化后，后台理解内容并关联旧记忆；更正与合并由用户审批。')
        for i,(t,b) in enumerate(NOTES[12]):
            panel(c,.8,2.65+i*1.35,4.6,1.17,'band');block(c,t,b,1.05,2.76+i*1.35,4.1,1.0,16)
        panel(c,5.75,2.6,6.75,4.14,'wing');shot(c,'dreaming',5.85,2.64,6.5,3.7,'图 1：后台整合资料并生成更正方案');floor(c)
    elif page==13:
        c.layout='bill total and evidence with lower value cards'
        c.text('434.50',.95,2.10,5.2,1.0,68,C,True)
        c.text('元 · 本例家庭账单合计',1.02,3.22,4.7,.42,22,G,True)
        c.text('电费 210.00    水费 68.50\n燃气 156.00',1.03,3.87,4.8,.8,20)
        shot(c,'bill',6.3,1.99,6.13,3.35,'图 1：助手在新会话中召回账单明细与来源')
        steps(c,[('保存清单','用户交付一次资料'),('新会话提问','召回明细并计算合计'),('核对来源','减少翻找与逐项抄录')],.85,6.02,11.6,.85)
        floor(c,7.02)
    elif page==15:
        c.layout='three input paths above software settings'
        steps(c,NOTES[15],.8,2.05,11.7,1.3)
        shot(c,'directory',.9,3.68,6.2,2.68,'图 1：用户选择授权目录并配置采集范围')
        panel(c,7.4,3.74,5.0,2.5,'wing')
        c.text('统一记忆管线',7.85,4.03,4.1,.43,27,C,True)
        c.text('授权内容 → 清洗与标准化\n质量校验 → 保留来源与范围',7.85,4.84,4.05,1.0,20)
        floor(c)
    elif page==16:
        c.layout='knowledge categories and provenance chain'
        for i,(t,b) in enumerate([('事实 FACT','时间 · 地点 · 金额'),('流程 WORKFLOW','先做什么 · 后做什么'),('案例 CASE','问题 · 处理 · 结果'),('模板 TEMPLATE','可以复用的结构')]):
            card(c,t,b,.8+i%2*2.75,2.25+i//2*1.45,2.5,1.17)
        shot(c,'source',6.54,2.1,5.96,3.45,'图 1：用户查看知识所关联的原始资料')
        c.text('知识不仅保存结论，也关联原文与文档块。\n点击来源即可核对依据。',.98,5.22,5.1,.78,17)
        panel(c,.8,6.24,11.7,.6,'band')
        c.text('来源证据  →  知识的状态、范围与版本  →  实体、关系与检索索引',1.05,6.39,11.2,.35,20,C,True,True);floor(c)
    elif page==17:
        c.layout='three retrieval lanes feeding result'
        steps(c,NOTES[17],.8,2.05,11.7,1.27)
        panel(c,.8,3.65,5.4,2.66,'wing')
        c.text('RRF 融合 → 重排',1.18,4.02,4.6,.48,29,C,True)
        c.text('范围、时间与有效状态过滤\n保留相关结果并附上来源',1.18,4.87,4.55,.95,20)
        shot(c,'keyword',6.55,3.51,5.8,2.86,'图 1：用户通过关键词检索知识并查看来源');floor(c)
    elif page==18:
        c.layout='ascending memory timeline'
        shot(c,'stage',.8,2.14,6.4,3.89,'图 1：用户查看短期记录与阶段记忆')
        for i,(t,b) in enumerate(NOTES[18]):
            x=7.35+i*.28;y=2.2+i*1.45
            panel(c,x,y,4.5,1.23,'band');block(c,t,b,x+.2,y+.1,4.1,1.06,16)
        c.text('任务上下文 → 阶段内容 → 持久知识',.9,6.57,11.5,.38,22,C,True,True);floor(c)
    elif page==19:
        c.layout='two conflict layers above approval evidence'
        steps(c,NOTES[19][:2],.8,2.05,11.7,1.37)
        shot(c,'dreaming',.85,3.63,6.4,2.87,'图 1：用户对照新旧内容并审批更正')
        panel(c,7.53,3.91,4.8,2.15,'round');block(c,'更正计划先审批','冻结目标版本\n用户核对后保存\n正文、版本与来源',7.88,4.22,4.05,1.62,20);floor(c)
    elif page==20:
        c.layout='online and reconnect swimlanes'
        intro(c,'共享记忆向在线节点传播；离线节点重连时发现缺失并补齐操作。')
        for i,(t,b) in enumerate(NOTES[20][:2]):
            panel(c,.8,2.83+i*1.69,5.1,1.43,'band');block(c,t,b,1.1,3.0+i*1.69,4.5,1.15,19)
        shot(c,'devices',6.3,2.69,6.1,3.6,'图 1：用户管理可信设备与同步设置')
        c.text('重建索引后继续本地检索，设备列表不等于实时送达证明',.9,6.57,11.6,.36,18,center=True);floor(c)
    elif page==21:
        c.layout='security quadrants with permissions window'
        for i,(t,b) in enumerate(NOTES[21]):
            x=.85+(i%2)*6.0;y=2.18+(i//2)*2.37
            card(c,t,b,x,y,5.58,1.78)
        shot(c,'settings',7.07,4.19,5.2,2.28,'图 1：用户设置助手记忆权限与共享范围');floor(c)
    elif page==22:
        c.layout='native platform bands and recovery metric'
        for i,t in enumerate(['银河麒麟 V11','Debian 兼容','签名升级']):
            panel(c,.85+i*4,2.04,3.62,.6,'band');c.text(t,1.05+i*4,2.17,3.22,.36,22,C,True,True)
        shot(c,'service',.9,2.97,7.0,3.51,'图 1：系统展示当前版本与原生能力状态')
        c.text('47',8.47,2.98,3.7,1.0,76,C,True)
        c.text('条记忆',10.35,3.60,1.7,.42,23)
        c.text('升级与恢复前后摘要一致',8.47,4.43,3.8,.73,24,G,True)
        c.text('麒麟优先使用系统能力\nDebian 支持基本记忆读写与检索',8.47,5.49,3.83,1.05,17);floor(c)
    elif page==23:
        c.layout='four-step lifecycle with reverse evidence placement'
        steps(c,[('任务开始','召回有效记忆'),('上下文注入','范围、来源与预算'),('规划与工具','工具执行与审批'),('结果沉淀','对话、工具与阶段内容')],.8,2.04,11.7,1.13)
        block(c,'从任务到可复用记忆','工具核算并写入文件后\n保存结果供后续会话使用',1.05,3.67,4.8,1.35,19)
        block(c,'上下文注入边界','按预算注入实际使用的来源\n记忆文本不升级为系统指令',1.05,5.16,4.8,1.35,19)
        shot(c,'task',6.55,3.45,5.8,3.02,'图 1：助手完成工具任务并保存处理结果');floor(c)
    elif page==25:
        c.layout='activity facts above process with full source'
        panel(c,.8,2.12,5.0,4.63,'wing')
        c.text('11 月 8 日',1.12,2.4,4.3,.62,36,C,True)
        c.text('19:00 · 社区天文台三层',1.12,3.22,4.25,.4,22)
        c.text('器材预算 2680 元',1.12,3.88,4.25,.4,23,G,True)
        c.line(1.12,4.57,5.43,4.57)
        block(c,'保留可复用流程','领取手册 → 检查器材\n分组记录 → 结束归还',1.12,4.95,4.25,1.45,20)
        shot(c,'source',6.06,2.29,6.35,3.84,'图 1：助手保留活动知识及其原始资料');floor(c)
    elif page==26:
        c.layout='cross-session editorial statement'
        c.text('换个会话\n继续问',.97,2.12,5.1,1.7,43,C,True)
        panel(c,.85,4.05,5.1,2.41,'round')
        block(c,'不用再交代背景','直接询问已保存的活动安排\n无需再次上传同一份资料',1.1,4.3,4.6,1.20,18)
        c.text('时间、地点与预算均附来源',1.1,5.95,4.6,.39,19,G,True)
        shot(c,'recall',6.35,2.17,6.1,4.02,'图 1：助手在新会话中召回已保存的活动资料');floor(c)
    elif page==27:
        c.layout='approval comparison with change callout'
        shot(c,'dreaming',.78,2.06,7.17,4.08,'图 1：用户核对活动安排的变化并批准更正')
        panel(c,8.35,2.28,4.02,3.92,'wing')
        c.text('11 月 8 日',8.73,2.60,3.3,.55,30)
        c.text('↓',9.8,3.27,.7,.5,30,C,True)
        c.text('11 月 15 日',8.73,3.91,3.3,.55,30,G,True)
        c.text('三层 → 二层\n预算与准备流程保留',8.73,4.86,3.3,.93,19)
        c.text('用户核对后批准更正，也可以保留原内容',.9,6.56,11.6,.37,21,C,True,True);floor(c)
    elif page==28:
        c.layout='updated activity outcome with retained budget'
        panel(c,.82,2.14,5.1,2.52,'round')
        c.text('11 / 15',1.14,2.43,4.5,.82,57,C,True)
        c.text('19:30 · 社区天文台二层',1.14,3.60,4.5,.47,22,G,True)
        panel(c,.82,5.01,5.1,1.53,'band')
        block(c,'预算仍为 2680 元','保留未变化的事实\n回答引用旧安排与更新通知',1.12,5.22,4.5,1.14,18)
        shot(c,'updated',6.23,2.19,6.2,4.02,'图 1：助手在后续查询中使用更正后的安排');floor(c)
    else:raise ValueError(page)
    return c

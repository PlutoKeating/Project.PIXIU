"""Verify latest user artwork preservation, screenshot authenticity and layout bounds."""
from pathlib import Path
import hashlib,json,zipfile,subprocess,re
from lxml import etree as ET
from PIL import Image,ImageDraw
from pptx import Presentation
WORK=Path(__file__).resolve().parents[1];ROOT=WORK.parents[1];OUT=WORK/'render/abc-trial'
m=json.loads((WORK/'review/abc-trial-manifest.json').read_text());rev=m['layout_revision']
ppt=OUT/'PIXIU项目报告-科技风试作版.pptx';source=ROOT/rev['source']
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
assert digest(ppt)==m['output_sha256'] and digest(source)==rev['source_sha256']
for i in m['inputs']:assert digest(ROOT/i['path'])==i['sha256'],i['path']
ns={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
changed=set()
for n in rev['modified_pages']:
 part=rev['parts'][str(n)];changed|={part,str(Path(part).parent/'_rels'/(Path(part).name+'.rels'))}
with zipfile.ZipFile(source) as before,zipfile.ZipFile(ppt) as after:
 assert after.testzip() is None
 for name in before.namelist():
  if name not in changed:assert before.read(name)==after.read(name),name
 for name in after.namelist():
  if name.endswith(('.xml','.rels')):ET.fromstring(after.read(name))
 for n in rev['modified_pages']:
  part=rev['parts'][str(n)];old=ET.fromstring(before.read(part));new=ET.fromstring(after.read(part))
  removed=m['slides'][n-1]['product_evidence']['removed_shape_ids']
  def keyed(root):return {s.xpath('.//p:cNvPr/@id',namespaces=ns)[0]:s for s in root.find('p:cSld/p:spTree',ns) if s.xpath('.//p:cNvPr/@id',namespaces=ns)}
  old,new=keyed(old),keyed(new)
  for key,shape in old.items():
   if key not in removed:assert ET.tostring(shape,method='c14n')==ET.tostring(new[key],method='c14n'),(n,key)
prs=Presentation(ppt);assert len(prs.slides)==34
captions=[];shots=0
for i,slide in enumerate(prs.slides,1):
 shapes={s.shape_id:s for s in slide.shapes}
 if i in rev['modified_pages']:
  evidence=m['slides'][i-1]['product_evidence'];assert evidence['artworks']
  for im in evidence['screenshots']:
   s=shapes[im['native_shape_id']];assert hashlib.sha256(s.image.blob).hexdigest()==im['sha256']
   x,y,w,h=im['box_inches'];assert x>=0 and y>=1.85 and x+w<=13 and y+h<=7.1
   iw,ih=im['native_size'];l,t,cw,ch=im['crop_pixels'];assert abs(s.width/s.height-cw/ch)<.001
   shots+=1
 for s in slide.shapes:
  if s.has_text_frame and s.text.startswith('图 '):
   t=s.text;captions.append(t);assert not any(v in t for v in ['（','）','(',')','红框','\n']);assert t.count('，')+t.count(',')<=1
assert len(captions)==25
# The 25th source image is retained verbatim on the unchanged chapter page 24.
assert shots==24
baseline='4bf08a7c3b504341cf4082b9f98f63fd2db98a521356674c1eecb36d530eef8e'
assert digest(WORK/'render/项目报告.pptx')==baseline and digest(ROOT/'docs/delivery/assets/项目报告.pptx')==baseline
pdf=OUT/'PIXIU项目报告-科技风试作版.pdf';info=subprocess.check_output(['pdfinfo',str(pdf)],text=True);assert re.search(r'Pages:\s+34\b',info)
overview=Image.new('RGB',(1600,2250),'#041F3B');d=ImageDraw.Draw(overview)
for i in range(34):
 im=Image.open(OUT/f'slide-{i+1:02}.png');assert im.size==(1600,900);im.load();im.thumbnail((400,225));x=i%4*400;y=i//4*250;overview.paste(im,(x,y));d.text((x+10,y+230),str(i+1),fill='white')
overview.save(OUT/'overview.png')
r={'pages':34,'captions':25,'redesigned_bodies':19,'source_sha256':digest(source),'pptx_sha256':digest(ppt),'pdf_sha256':digest(pdf),'checks':['Latest cover and closing bytes preserved','All unmodified package parts preserved','Headers and folios preserved','Original screenshot bytes and aspect ratios verified','Native ABC artwork recorded per slide','Single-sentence captions','34 rendered pages','Formal assets unchanged']}
(WORK/'review/abc-trial-validation.json').write_text(json.dumps(r,ensure_ascii=False,indent=2)+'\n');print(json.dumps(r,ensure_ascii=False))

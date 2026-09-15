from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from lxml import etree as E
from copy import deepcopy
from PIL import Image,ImageChops
import zipfile,json,hashlib,re
root=Path(__file__).resolve().parents[3];work=root/'submission/presentation-production'
import argparse
parser=argparse.ArgumentParser();parser.add_argument('--rebuilt',type=Path,required=True);args=parser.parse_args()
a=work/'source/user-submission-20260915-36pages.pptx';b=next((root/'submission').glob('603821*/*/项目报告.pptx'))
x,y=Presentation(a),Presentation(b)
assert len(x.slides)==len(y.slides)==36
assert len(x.slides[5].shapes)==len(y.slides[5].shapes)==60
for s,t in zip(x.slides[5].shapes,y.slides[5].shapes):
 assert (s.shape_id,s.left,s.top,s.width,s.height)==(t.shape_id,t.left,t.top,t.width,t.height)
 assert E.tostring(s._element.find('{*}spPr'))==E.tostring(t._element.find('{*}spPr')) if s._element.find('{*}spPr') is not None else True
for n,(s,t) in enumerate(zip(x.slides,y.slides),1):
 if n in [6,12,25]:continue
 assert len(s.shapes)==len(t.shapes)
 for ss,tt in zip(s.shapes,t.shapes):
  if ss.has_text_frame and ss.left>Inches(11.5) and ss.top>Inches(7) and ss.text.strip().isdigit():
   assert tt.text==f'{n:02}';continue
  assert E.tostring(ss._element)==E.tostring(tt._element),(n,ss.name)
for idx in [24,31,38]:
 assert hashlib.sha256(x.slides[11].shapes[idx].image.blob).digest()==hashlib.sha256(y.slides[11].shapes[idx].image.blob).digest()
assert E.tostring(x.slides[24].shapes[23]._element)==E.tostring(y.slides[24].shapes[23]._element)
all_text=['\n'.join(t.rstrip() for t in s._element.xpath('.//a:t/text()')) for s in y.slides]
assert 'XYZ' not in all_text[5]
for n in range(2,36):
 assert f'{n:02}' in y.slides[n-1]._element.xpath('.//a:t/text()'),n
with zipfile.ZipFile(a) as z:old={n:z.read(n) for n in z.namelist()}
with zipfile.ZipFile(b) as z:new={n:z.read(n) for n in z.namelist()};assert z.testzip() is None
with zipfile.ZipFile(args.rebuilt) as z:rebuilt={n:z.read(n) for n in z.namelist()}
assert new==rebuilt
changed=[n for n in old if old[n]!=new[n]]
assert all(re.fullmatch(r'ppt/slides/slide\d+.xml',n) or re.fullmatch(r'ppt/slides/_rels/slide\d+.xml.rels',n) for n in changed)
assert set(new)-set(old)=={'ppt/media/pixiu-policy-01.png','ppt/media/pixiu-policy-02.png','ppt/media/pixiu-policy-03.png'}
report={'slides':36,'source_sha256':hashlib.sha256(a.read_bytes()).hexdigest(),'output_sha256':hashlib.sha256(b.read_bytes()).hexdigest(),'policy_shape_count':60,'policy_all_shape_ids_geometry_and_style_unchanged':True,'all_other_slide_shapes_unchanged_except_footer':True,'three_device_screenshot_bytes_unchanged':True,'bibliography_unchanged':True,'visible_page_numbers':'02–35; covers 1/36 intentionally unnumbered','rebuilt_package_parts_identical':True,'changed_existing_parts':changed,'added_parts':sorted(set(new)-set(old))}
(work/'review/submission-36page-validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n')
(work/'review/submission-36page-text.md').write_text('# 当前交付版36页完整文本\n\n封面与尾页为原始图片；其余文本包括分组与表格。\n\n'+'\n\n'.join(f'## 第 {i} 页\n\n{t.strip()}' for i,t in enumerate(all_text,1)))
print(json.dumps(report,ensure_ascii=False))

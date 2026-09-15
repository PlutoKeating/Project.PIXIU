"""Focused native slide edits; reuse the existing chapter-cover artwork."""
from copy import deepcopy
from lxml import etree as ET
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

NS = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
EMU=914400
C='57F1FF'; WHITE='FFFFFF'; GOLD='FFDCAB'
COVERS={33:('01','场景需求与使用背景','从多设备使用中的记忆断点，理解用户需求',['使用场景','需求回应']),
        34:('02','产品定位与方案架构','面向 OS Agent，建立持续可用的记忆能力',['产品定位','技术全景']),
        35:('03','两大核心功能亮点','让记忆跨设备接续，并随新资料持续整合',['两大亮点','记忆共享','自动记忆','场景示例']),
        36:('06','服务模式与商业探索','围绕部署、持续维护与生态合作，探索服务路径',['服务模式','试点路线'])}


def text(n):return ''.join(n.xpath('.//a:t/text()',namespaces=NS))
def off(n):
    nodes=n.xpath('./p:spPr/a:xfrm/a:off',namespaces=NS)
    return nodes[0] if nodes else None

def set_text(n,value):
    runs=n.xpath('.//a:t',namespaces=NS)
    assert runs
    runs[0].text=value
    for r in runs[1:]:r.text=''


def cover(root, number):
    """Copy the technical divider's full composition and native effects."""
    num,title,subtitle,labels=COVERS[number]
    tree=root.find('p:cSld/p:spTree',NS)
    for n in list(tree):
        value=text(n)
        if value=='04':set_text(n,num)
        elif value=='持续记忆技术体系':set_text(n,title)
        elif value=='数据有结构，更新有依据，协作有边界':set_text(n,subtitle)
        elif value=='12':set_text(n,str(number))
        o=off(n)
        if o is None:continue
        x,y=int(o.get('x'))/EMU,int(o.get('y'))/EMU
        if 5.7<y<6.3:
            slot=round((x-.92)/2.42)
            if slot>=len(labels):tree.remove(n);continue
            # Preserve each native capsule and text styling, distribute pairs evenly.
            width=12/len(labels)-.30
            target=.92+slot*(12/len(labels))
            is_text=bool(value)
            o.set('x',str(round((target+(.07 if is_text else 0))*EMU)))
            n.xpath('./p:spPr/a:xfrm/a:ext',namespaces=NS)[0].set('cx',str(round((width-(.14 if is_text else 0))*EMU)))
            if is_text:set_text(n,labels[slot])
    return root


def explain(root, kind):
    tree=root.find('p:cSld/p:spTree',NS)
    temp=Presentation(); s=temp.slides.add_slide(temp.slide_layouts[6])
    def add(t,x,y,w,h,size=18,bold=False,color=WHITE):
        b=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        tf=b.text_frame;tf.word_wrap=True
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for i,line in enumerate(t.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
            p.text=line;p.font.name='Microsoft YaHei';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color)
            p.line_spacing=1.08;p.space_after=Pt(0)
        return b
    def card(x,y,w,h):
        a=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string('08456A');a.line.color.rgb=RGBColor.from_string('237F9F');a.line.width=Pt(.6);a.adjustments[0]=.08
    for n in list(tree):
        o=off(n)
        if o is None:continue
        x,y=int(o.get('x'))/EMU,int(o.get('y'))/EMU
        if y<1.85 or y>7.1:continue
        value=text(n)
        if kind=='sharing':
            if value in {'分布式多设备 Agent 记忆共享','活动安排','资料来源','时间 / 地点','原文 / 版本'}:tree.remove(n);continue
            # Complete peer diagram remains native and connected on the left.
            scale=.56
            o.set('x',str(round((.72+(x-.70)*scale)*EMU)))
            o.set('y',str(round((3.10+(y-1.92)*scale)*EMU)))
            for ext in n.xpath('./p:spPr/a:xfrm/a:ext',namespaces=NS):
                for attr in ['cx','cy']:ext.set(attr,str(round(int(ext.get(attr))*scale)))
            for r in n.xpath('.//a:rPr | .//a:defRPr | .//a:endParaRPr',namespaces=NS):
                if r.get('sz'):r.set('sz',str(round(int(r.get('sz'))*scale)))
        else:
            # Retain the actual screenshot and its frame; replace the large totals.
            if x>=7.8:tree.remove(n);continue
            o.set('x',str(round((.95+(x-.65)*.78)*EMU)))
            o.set('y',str(round((2.72+(y-1.93)*.78)*EMU)))
            for ext in n.xpath('./p:spPr/a:xfrm/a:ext',namespaces=NS):
                ext.set('cx',str(round(int(ext.get('cx'))*.78)))
                ext.set('cy',str(round(int(ext.get('cy'))*.78)))
    if kind=='sharing':
        add('在书房保存的活动安排，换到客厅或笔记本，Agent 仍能查到先前的内容与来源。\nPIXIU 将授权共享的记忆在可信设备间同步，各端保留本地副本。',.78,1.98,11.85,.83,19)
        for y,title,body in [(3.02,'01  用户只需建立信任','配对设备，选择共享范围。\n私人记忆保持私有，共享由用户决定。'),(4.33,'02  PIXIU 负责接续记忆','在对等设备间同步知识、来源与更新。\n断连保留本地副本，重连后补齐差异。'),(5.64,'03  换设备也能继续做事','新会话可以继续查询、核对和复用。\n减少重复交代背景，也能追溯原始依据。')]:
            card(7.68,y,4.89,1.16);add(title,7.9,y+.13,4.45,.31,18,True,GOLD);add(body,7.9,y+.51,4.45,.53,15)
        for x,y in [(4.98,1.92),(.85,4.28),(9.12,4.28)]:
            add('共享知识与来源',.72+(x-.70)*.56+.16,3.10+(y-1.92)*.56+.80,1.58,.22,9)
        add('三台设备对等互连 · 共享范围由用户控制',.91,6.29,6.40,.43,17,True,C)
    else:
        add('家庭开销核对：清单曾经交给助手，过后只记得“水电燃气花了钱”。',.76,1.96,11.86,.48,19)
        for y,title,body in [(2.72,'用户做什么','先把清单交给助手保存；在新会话中\n询问九月明细、合计，并要求提供来源。'),(4.02,'PIXIU 实现什么','召回之前保存的账单记录，列出明细、\n核算合计，并保留可回看的原始依据。'),(5.32,'用户得到什么','得到可核对的答案，无需重新翻文件、\n逐项抄录和相加。')]:
            card(7.56,y,5.05,1.18);add(title,7.8,y+.12,4.57,.31,18,True,GOLD);add(body,7.8,y+.50,4.57,.54,15)
        add('本例结果：210.00 + 68.50 + 156.00 = 434.50 元',.77,6.64,11.75,.36,19,True,C)
    next_id=max(int(x) for x in root.xpath('.//p:cNvPr/@id',namespaces=NS))+1
    for a in s.shapes:
        node=ET.fromstring(ET.tostring(a._element))
        for identity in node.xpath('.//p:cNvPr',namespaces=NS):identity.set('id',str(next_id));next_id+=1
        tree.append(node)
    return root


def align_knowledge(root):
    """Align the four categories and the source/knowledge/index flow to a grid."""
    tree=root.find('p:cSld/p:spTree',NS)
    nodes=list(tree)
    # spTree has two non-shape children before the authored shapes.
    shapes=[n for n in nodes if n.tag.rsplit('}',1)[-1] in {'sp','pic','cxnSp','grpSp','graphicFrame'}]
    def box(index,x,y,w,h):
        node=shapes[index]
        transform=node.xpath('./p:spPr/a:xfrm',namespaces=NS)[0]
        transform.find('a:off',NS).set('x',str(round(x*EMU)))
        transform.find('a:off',NS).set('y',str(round(y*EMU)))
        transform.find('a:ext',NS).set('cx',str(round(w*EMU)))
        transform.find('a:ext',NS).set('cy',str(round(h*EMU)))
    for i in range(4):
        x=.85+3.0*i;start=7+4*i
        box(start,x,2.20,2.70,1.25)
        box(start+1,x+.945,2.20,.81,0)
        box(start+2,x+.10,2.43,2.50,.35)
        box(start+3,x+.10,3.00,2.50,.32)
    box(26,.85,4.40,3.10,1.80)
    box(27,1.935,4.40,.93,0)
    box(28,1.00,4.78,2.80,.46)
    box(29,1.00,5.49,2.80,.38)
    box(23,5.10,4.40,3.10,1.80)
    box(24,5.25,4.78,2.80,.46)
    box(25,5.25,5.49,2.80,.38)
    box(30,4.08,5.30,.87,0)
    for i,(card,label) in enumerate([(32,33),(35,36),(38,39)]):
        y=4.35+.72*i
        box(card,9.65,y,2.90,.46)
        box(label,9.72,y+.07,2.76,.32)
    # Native right-angle distribution replaces three diagonal fans.
    template=deepcopy(shapes[34])
    for index in [31,34,37]:tree.remove(shapes[index])
    next_id=max(int(v) for v in root.xpath('.//p:cNvPr/@id',namespaces=NS))+1
    for x,y,w,h,arrow in [(8.32,5.30,.58,0,False),(8.90,4.58,0,1.44,False),
                          (8.90,4.58,.61,0,True),(8.90,5.30,.61,0,True),(8.90,6.02,.61,0,True)]:
        node=deepcopy(template)
        xf=node.xpath('./p:spPr/a:xfrm',namespaces=NS)[0]
        xf.attrib.pop('flipH',None);xf.attrib.pop('flipV',None)
        xf.find('a:off',NS).set('x',str(round(x*EMU)));xf.find('a:off',NS).set('y',str(round(y*EMU)))
        xf.find('a:ext',NS).set('cx',str(round(w*EMU)));xf.find('a:ext',NS).set('cy',str(round(h*EMU)))
        if not arrow:
            for end in node.xpath('.//a:tailEnd | .//a:headEnd',namespaces=NS):end.set('type','none')
        for identity in node.xpath('.//p:cNvPr',namespaces=NS):identity.set('id',str(next_id));next_id+=1
        tree.append(node)
    return root

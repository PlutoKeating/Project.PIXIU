"""Add short judge-facing introductions to the three illustrated logic slides.

Use the archived latest deck as input; output must be a new candidate path.
Only slide XML is replaced, retaining screenshots and every other package part.
"""
import argparse,zipfile
from pathlib import Path
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.enum.text import PP_ALIGN,MSO_ANCHOR
from pptx.dml.color import RGBColor
COPY={
19:'PIXIU 按使用阶段管理 Agent 记忆：短期保留当前任务要点，中期保存阶段内容。\n用户选择的内容进入长期知识库，并在后续任务中通过检索召回复用。',
20:'多设备编辑既可能产生同一知识的版本差异，也可能出现内容矛盾。\nPIXIU 分别用副本收敛与业务仲裁处理；涉及更正、合并的方案由用户核对后保存。',
21:'可信设备按授权范围共享记忆：在线时传播签名操作，并在对端保存为本地副本。\n离线时仍可使用本地记忆；重连后交换摘要、补齐缺失操作，使副本逐步一致。'}
def resize(q,scale,cx,oldy,newy):
 q.left=round(cx+(q.left-cx)*scale);q.top=round(newy+(q.top-oldy)*scale);q.width=round(q.width*scale);q.height=round(q.height*scale)
 if q.has_text_frame:
  for p in q.text_frame.paragraphs:
   for r in p.runs:
    if r.font.size:r.font.size=round(r.font.size*scale)
def intro(s,num,q=None):
 if q is None:q=s.shapes.add_textbox(Inches(.80),Inches(2.04),Inches(11.73),Inches(.52))
 q.left=Inches(.80);q.top=Inches(2.04);q.width=Inches(11.73);q.height=Inches(.52)
 f=q.text_frame;f.clear();f.word_wrap=False;f.vertical_anchor=MSO_ANCHOR.MIDDLE;f.margin_left=f.margin_right=f.margin_top=f.margin_bottom=0
 for i,line in enumerate(COPY[num].split('\n')):
  p=f.paragraphs[0] if i==0 else f.add_paragraph();p.alignment=PP_ALIGN.CENTER;p.space_before=p.space_after=0;p.line_spacing=Pt(21)
  r=p.add_run();r.text=line;r.font.name='Microsoft YaHei';r.font.size=Pt(16.5);r.font.color.rgb=RGBColor.from_string('EAF8FF')
 return q
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args();assert not a.output.exists()
 r=Presentation(a.source);assert len(r.slides)==36
 # Memory: uniformly scale the illustration region, maintaining screenshot aspect ratio.
 s=r.slides[18]
 for q in list(s.shapes):
  if q.top>=Inches(2.0) and q.top<Inches(6.55):resize(q,.84,Inches(6.6667),Inches(2.14),Inches(2.77))
 intro(s,19)
 # Conflict: retain small diagram labels at full size; shift the two rails together.
 s=r.slides[19]
 for q in list(s.shapes):
  if Inches(2.0)<=q.top<Inches(3.45):q.top+=Inches(.60)
  elif q.shape_id in [83,84,85]:resize(q,.82,Inches(4.05),Inches(3.63),Inches(4.15))
  elif q.shape_id in [87,88,89]:q.top+=Inches(.40)
 intro(s,20)
 # Sync already reserves an introduction band; expand it to two concise sentences.
 s=r.slides[20];q=next(q for q in s.shapes if q.shape_id==78);intro(s,21,q)
 with zipfile.ZipFile(a.source) as z:entries=z.infolist();blobs={i.filename:z.read(i.filename) for i in entries}
 for n in COPY:blobs[f'ppt/slides/slide{n}.xml']=E.tostring(r.slides[n-1]._element,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(a.output,'w') as z:
  for i in entries:z.writestr(i,blobs[i.filename])
 print(a.output)

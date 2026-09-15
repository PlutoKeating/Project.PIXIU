"""Complete chapter 04 navigation using the latest deck's native styles."""
import argparse,zipfile
from pathlib import Path
from copy import deepcopy
from lxml import etree as E
from pptx import Presentation
from pptx.util import Inches
from pptx.shapes.autoshape import Shape
NS={'p':'http://schemas.openxmlformats.org/presentationml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
LABELS=['多源接入','知识检索','记忆流转','对等同步','安全部署','任务循环','技术背书']
ACTIVE={16:0,17:1,18:1,19:2,20:2,21:3,22:4,23:4,24:5,25:6}
def clone(s,e,x,y,w,h,t=None):
 e=deepcopy(e);id=max(q.shape_id for q in s.shapes)+1
 c=e.find('.//p:cNvPr',NS);c.set('id',str(id));c.set('name','Chapter navigation '+str(id))
 s.shapes._spTree.insert_element_before(e,'p:extLst');q=Shape(e,s.shapes)
 q.left=Inches(x);q.top=Inches(y);q.width=Inches(w);q.height=Inches(h)
 if t is not None:
  texts=e.findall('.//a:t',NS);assert texts;texts[0].text=t
  for a in texts[1:]:a.text=''
 return q
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--source',type=Path,required=True);p.add_argument('--output',type=Path,required=True);a=p.parse_args();assert not a.output.exists()
 r=Presentation(a.source);assert len(r.slides)==36
 ref={q.shape_id:deepcopy(q._element) for q in r.slides[24].shapes}
 assert '研究' in ''.join(q.text for q in r.slides[24].shapes if q.has_text_frame)
 width=(9.54-.10*6)/7;step=width+.10
 for n,active in ACTIVE.items():
  s=r.slides[n-1]
  assert any(q.has_text_frame and q.text=='技术实现' for q in s.shapes)
  for q in list(s.shapes):
   if Inches(2.70)<q.left and Inches(.20)<=q.top<Inches(.80):s.shapes._spTree.remove(q._element)
  for i,t in enumerate(LABELS):
   x=2.78+i*step
   clone(s,ref[63],x,.22,width,.42)
   if i==active:
    clone(s,ref[74],x+.02,.24,width-.04,.38)
    clone(s,ref[75],x+.19,.66,width-.38,0)
   clone(s,ref[76 if i==active else 64],x+.03,.29,width-.06,.29,t)
 s=r.slides[14];assert any(q.has_text_frame and q.text=='技术架构与实现方案' for q in s.shapes)
 old={q.shape_id:deepcopy(q._element) for q in s.shapes}
 for q in list(s.shapes):
  if q.shape_id in range(9,19):s.shapes._spTree.remove(q._element)
 width=(11.81-.14*6)/7
 for i,t in enumerate(LABELS):
  x=.92+i*(width+.14)
  clone(s,old[9],x,5.75,width,.46)
  clone(s,old[10],x+.07,5.82,width-.14,.338,t)
 with zipfile.ZipFile(a.source) as z:entries=z.infolist();blobs={i.filename:z.read(i.filename) for i in entries}
 for n in range(15,26):blobs[f'ppt/slides/slide{n}.xml']=E.tostring(r.slides[n-1]._element,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(a.output,'w') as z:
  for i in entries:z.writestr(i,blobs[i.filename])
 print(a.output)

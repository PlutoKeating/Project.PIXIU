"""Supplement spreadsheet cells with formulas, annotations and embedded visuals.

openpyxl owns OOXML relationships and drawing decoding. No workbook is executed,
recalculated, saved back, or allowed to follow external links.
"""
import io
import warnings

from openpyxl import load_workbook
from openpyxl.xml.functions import tostring


def supplement_workbook(document, data):
    with warnings.catch_warnings(record=True) as notices:
        warnings.simplefilter("always", UserWarning)
        workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
        try:
            for sheet in [*workbook.worksheets, *workbook.chartsheets]:
                # Chartsheets have drawings but no cells.
                if hasattr(sheet, "iter_rows"):
                    details = []
                    for row in sheet.iter_rows():
                        for cell in row:
                            location = f"{sheet.title}!{cell.coordinate}"
                            if cell.data_type == "f":
                                formula = cell.value if isinstance(cell.value, str) else getattr(cell.value, "text", "")
                                details.append(f"{location} 公式（未重新计算）：{formula}")
                            if cell.comment:
                                details.append(f"{location} 批注：{cell.comment.text}")
                            if cell.hyperlink:
                                details.append(f"{location} 链接（未访问）：{cell.hyperlink.target or cell.hyperlink.location}")
                    document.text(f"工作表 {sheet.title} 补充内容", "\n".join(details))
                # Pinned openpyxl 3.1.5 drawing collections. Keep this dependency
                # detail inside the format adapter, outside the MCP contract.
                for index, picture in enumerate(sheet._images):
                    anchor = getattr(picture.anchor, "_from", None)
                    position = f"R{anchor.row + 1}C{anchor.col + 1}" if anchor else str(index + 1)
                    document.image(f"工作表 {sheet.title} 图片 {position}", picture._data())
                for index, chart in enumerate(sheet._charts):
                    document.text(f"工作表 {sheet.title} 图表 {index + 1} 结构与数据引用",
                                  tostring(chart._write()).decode("utf-8"))
        finally:
            workbook.close()
    document.warnings.extend(f"工作簿内容未完整解析：{notice.message}" for notice in notices)
